"""
Générateur des modèles Excel téléchargeables pour l'importation
massive (module 15, section 9). Pure présentation : aucune formule de
paie, aucune logique métier — uniquement des exemples illustratifs.
"""

from pathlib import Path

from openpyxl import Workbook
from exports.logo_excel import ajouter_logo
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from models.enums import TypeImport

POLICE_ENTETE = Font(name="Arial", size=10, bold=True, color="FFFFFF")
REMPLISSAGE_ENTETE = PatternFill("solid", fgColor="2F5496")

MODELES: dict = {
    TypeImport.ENSEIGNANTS: {
        "colonnes": ["Nom", "Prenom", "Sexe", "Statut", "Taux_Horaire", "Telephone", "Email", "Adresse"],
        "exemples": [
            ["Kamgang", "Jean Paul", "M", "P", 2000, "", "", ""],
            ["Ngono", "Marie", "F", "", "", "", "", ""],
        ],
        "instructions": [
            "Seul le Nom est obligatoire. Une colonne laissée vide (Sexe, Statut, Taux_Horaire) se complète "
            "plus tard dans Gestion › Enseignants ; la fiche reste alors hors de la paie.",
            "Sexe : M ou F (Masculin, Féminin, Homme, Femme sont aussi acceptés).",
            "Statut : P (Permanent) ou V (Vacataire) ; les mots entiers sont aussi acceptés.",
            "Taux_Horaire : nombre entier en FCFA (« 1500 », « 1 500 » ou « 1500 FCFA »).",
            "Une liste existante en Word ou en PDF peut aussi être importée directement, sans ce modèle.",
        ],
    },
    TypeImport.HEURES: {
        "colonnes": ["Nom", "Prenom", "Semaine_1", "Semaine_2", "Semaine_3", "Semaine_4", "Semaine_5"],
        "exemples": [["Kamgang", "Jean Paul", 20, 20, 20, 20, 20]],
        "instructions": [
            "Nom/Prenom doivent correspondre exactement à un enseignant déjà existant.",
            "Chaque colonne Semaine_N est optionnelle ; laisser vide si aucune heure.",
            "L'import n'est possible que pour une période au statut OUVERTE.",
        ],
    },
    TypeImport.REMUNERATIONS: {
        "colonnes": ["Nom", "Prenom", "Prime_AP_PP", "Surveillance_Secretariat", "Indemnite_Suggestion_Admin"],
        "exemples": [["Kamgang", "Jean Paul", 20000, 10000, 5000]],
        "instructions": [
            "Montants en FCFA entiers, sans décimales, jamais négatifs.",
            "L'import n'est possible que pour une période au statut OUVERTE.",
        ],
    },
    TypeImport.RETENUES: {
        "colonnes": ["Nom", "Prenom", "Retenue_Amicale", "Dette"],
        "exemples": [["Kamgang", "Jean Paul", 5000, 10000]],
        "instructions": [
            "Montants en FCFA entiers, sans décimales, jamais négatifs.",
            "L'import n'est possible que pour une période au statut OUVERTE.",
        ],
    },
}


def generer_classeur_modele(type_import: TypeImport) -> Workbook:
    """Construit le classeur modèle (colonnes + exemple + feuille Instructions) pour un type d'import donné."""
    modele = MODELES[type_import]
    classeur = Workbook()

    feuille = classeur.active
    feuille.title = "Donnees"
    for index, nom_colonne in enumerate(modele["colonnes"], start=1):
        cellule = feuille.cell(row=1, column=index, value=nom_colonne)
        cellule.font = POLICE_ENTETE
        cellule.fill = REMPLISSAGE_ENTETE
        feuille.column_dimensions[get_column_letter(index)].width = 20

    for ligne_index, exemple in enumerate(modele["exemples"], start=2):
        for colonne_index, valeur in enumerate(exemple, start=1):
            feuille.cell(row=ligne_index, column=colonne_index, value=valeur)

    feuille_instructions = classeur.create_sheet("Instructions")
    feuille_instructions.cell(row=1, column=1, value="Instructions de remplissage").font = Font(bold=True, size=12)
    for i, ligne_instruction in enumerate(modele["instructions"], start=3):
        feuille_instructions.cell(row=i, column=1, value=f"• {ligne_instruction}")
    feuille_instructions.column_dimensions["A"].width = 90

    return ajouter_logo(classeur)


def generer_fichier_modele(type_import: TypeImport, dossier: Path) -> Path:
    """Génère le fichier modèle sur disque : template_import_{type}.xlsx."""
    dossier.mkdir(parents=True, exist_ok=True)
    classeur = generer_classeur_modele(type_import)
    chemin = dossier / f"template_import_{type_import.value}.xlsx"
    classeur.save(chemin)
    return chemin


def generer_modele_csv(type_import: TypeImport) -> str:
    """Génère le contenu CSV du modèle (en-têtes + exemple), sous forme de texte prêt à télécharger."""
    modele = MODELES[type_import]
    lignes = [",".join(modele["colonnes"])]
    for exemple in modele["exemples"]:
        lignes.append(",".join(str(v) for v in exemple))
    return "\n".join(lignes)


REMPLISSAGE_A_COMPLETER = PatternFill("solid", fgColor="FFF2CC")


def generer_classeur_fiches_a_completer(enseignants) -> Workbook:
    """
    Liste des fiches à compléter, au format du modèle d'import Enseignants :
    les informations connues sont pré-remplies, les cases manquantes sont
    surlignées. Une fois remplie, elle se réimporte telle quelle
    (stratégie « Mettre à jour les doublons ») : seules les cases
    renseignées sont reprises, rien n'est effacé.
    """
    modele = MODELES[TypeImport.ENSEIGNANTS]
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Donnees"
    for index, nom_colonne in enumerate(modele["colonnes"], start=1):
        cellule = feuille.cell(row=1, column=index, value=nom_colonne)
        cellule.font = POLICE_ENTETE
        cellule.fill = REMPLISSAGE_ENTETE
        feuille.column_dimensions[get_column_letter(index)].width = 20
    feuille.freeze_panes = "A2"

    for ligne_index, enseignant in enumerate(enseignants, start=2):
        valeurs = [
            enseignant.nom, enseignant.prenom,
            enseignant.sexe.value if enseignant.sexe is not None else None,
            enseignant.statut.value if enseignant.statut is not None else None,
            enseignant.taux_horaire,
            enseignant.telephone, enseignant.email, enseignant.adresse,
        ]
        for colonne_index, valeur in enumerate(valeurs, start=1):
            cellule = feuille.cell(row=ligne_index, column=colonne_index, value=valeur)
            if colonne_index in (3, 4, 5) and valeur is None:
                cellule.fill = REMPLISSAGE_A_COMPLETER

    feuille_instructions = classeur.create_sheet("Instructions")
    feuille_instructions.cell(row=1, column=1, value="Fiches d'enseignants à compléter").font = Font(bold=True, size=12)
    instructions = [
        "Remplissez les cases surlignées en jaune (Sexe, Statut, Taux_Horaire). Une case encore inconnue peut "
        "rester vide : la fiche restera « à compléter ».",
        "Ne modifiez pas les colonnes Nom et Prenom : elles servent à retrouver chaque enseignant.",
        "Sexe : M ou F. Statut : P (Permanent) ou V (Vacataire). Taux_Horaire : nombre entier en FCFA.",
        "Réimportez ensuite ce fichier dans Documents & Opérations › Importation, type « Enseignants », en "
        "choisissant la stratégie « Mettre à jour les doublons ». Seules les cases renseignées sont reprises : "
        "aucune information déjà enregistrée n'est effacée.",
    ]
    for i, ligne_instruction in enumerate(instructions, start=3):
        feuille_instructions.cell(row=i, column=1, value=f"• {ligne_instruction}")
    feuille_instructions.column_dimensions["A"].width = 110
    return ajouter_logo(classeur)
