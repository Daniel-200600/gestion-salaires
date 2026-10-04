"""
Logo de l'établissement sur les classeurs Excel exportés.

Le logo (Administration › Paramètres › En-tête du bulletin et logo) est
placé en haut de chaque feuille, à droite du tableau : aucune cellule
n'est déplacée, les fichiers restent lisibles et réimportables tels quels.
Sans logo enregistré, le classeur est rendu inchangé.
"""

import io
import logging
from typing import Optional

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ImageExcel
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

HAUTEUR_LOGO_PX = 64


def logo_etablissement() -> Optional[bytes]:
    try:
        from services.identite_etablissement_service import obtenir_identite

        identite = obtenir_identite()
        return identite.logo if identite is not None else None
    except Exception as erreur:  # noqa: BLE001 — un export ne doit jamais échouer à cause du logo
        logger.warning("Logo de l'établissement non lu : %s", erreur)
        return None


def ajouter_logo(classeur: Workbook, logo: Optional[bytes] = None) -> Workbook:
    """Ajoute le logo en haut à droite de chaque feuille du classeur (si un logo est enregistré)."""
    logo = logo if logo is not None else logo_etablissement()
    if not logo:
        return classeur
    for feuille in classeur.worksheets:
        try:
            image = ImageExcel(io.BytesIO(logo))
        except Exception as erreur:  # noqa: BLE001 — image illisible : export sans logo
            logger.warning("Logo de l'établissement illisible : %s", erreur)
            return classeur
        ratio = HAUTEUR_LOGO_PX / image.height if image.height else 1
        image.height, image.width = HAUTEUR_LOGO_PX, int(image.width * ratio)
        feuille.add_image(image, f"{get_column_letter(feuille.max_column + 2)}1")
    return classeur
