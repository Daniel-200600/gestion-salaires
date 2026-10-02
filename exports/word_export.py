"""
Génération du fichier Word individuel (module 07).

Charge le modèle Word actif — par défaut templates/bulletin_template.docx
(créé UNE SEULE FOIS par templates/build_template.py), sinon un modèle
importé par l'administrateur (services/modele_bulletin_service.py) — et y injecte des valeurs déjà préparées
par services/bulletin_service.py (lui-même dérivé de
services/paie_service.py). Ce module ne contient et ne doit JAMAIS
contenir de formule de paie (gain heures, taxe, net à percevoir) :
il se limite au remplacement de texte et à l'enregistrement du
document — pure présentation d'une donnée déjà calculée.

Ce module ne dépend pas de Streamlit.
"""

from pathlib import Path
import io
import re
from dataclasses import dataclass
from typing import Dict, Iterator, List, Optional, Tuple

from docx import Document

from config.settings import DATA_DIR, chemin_modele_standard
from exports.excel_export import chemin_sortie_disponible
from utils.balises_bulletin import contient_balise_residuelle, remplacer_balises, trouver_balises, valeurs_par_nom
from utils.detection_bulletin import Segment, genre_texte, mois_annee
from utils.formatters import nettoyer_nom_fichier

TEMPLATE_PATH = chemin_modele_standard("bulletin_template.docx")
EXPORT_DIR_BULLETINS = DATA_DIR / "exports" / "bulletins"


class WordExportError(Exception):
    """Levée pour toute impossibilité de générer un bulletin Word (template introuvable, etc.)."""


def _parcourir(document) -> Iterator[Tuple[object, int, int]]:
    """
    Parcourt TOUS les paragraphes d'un document — corps, tableaux (y
    compris imbriqués), en-têtes et pieds de page — et fournit pour
    chacun un numéro de « ligne » (ligne de tableau ou paragraphe hors
    tableau) et un numéro de colonne (case dans la ligne). Ces repères
    servent à la détection des champs d'un bulletin déjà rempli.
    """
    compteur = [0]

    def _nouvelle_ligne() -> int:
        compteur[0] += 1
        return compteur[0]

    def _dans_tableaux(tables):
        for table in tables:
            for ligne in table.rows:
                rang = _nouvelle_ligne()
                vues = set()
                for colonne, cellule in enumerate(ligne.cells):
                    # Une cellule fusionnée apparaît plusieurs fois dans ligne.cells.
                    if id(cellule._tc) in vues:
                        continue
                    vues.add(id(cellule._tc))
                    for paragraphe in cellule.paragraphs:
                        yield paragraphe, rang, colonne
                    yield from _dans_tableaux(cellule.tables)

    for paragraphe in document.paragraphs:
        yield paragraphe, _nouvelle_ligne(), 0
    yield from _dans_tableaux(document.tables)
    for section in document.sections:
        for partie in (section.header, section.footer, section.first_page_header,
                       section.first_page_footer, section.even_page_header, section.even_page_footer):
            if partie is None or partie.is_linked_to_previous:
                continue
            for paragraphe in partie.paragraphs:
                yield paragraphe, _nouvelle_ligne(), 0
            yield from _dans_tableaux(partie.tables)


def iterer_paragraphes(document) -> Iterator:
    """Tous les paragraphes du document (voir _parcourir) : un modèle importé peut placer ses balises partout."""
    for paragraphe, _rang, _colonne in _parcourir(document):
        yield paragraphe


def texte_document(document) -> str:
    """Texte complet du document (tous paragraphes), une ligne par paragraphe."""
    return "\n".join(p.text for p in iterer_paragraphes(document))


def extraire_balises_document(document) -> List[str]:
    """Noms des balises présentes dans le document, dans l'ordre d'apparition (sans doublon)."""
    noms = []
    for paragraphe in iterer_paragraphes(document):
        for occurrence in trouver_balises(paragraphe.text):
            if occurrence.nom not in noms:
                noms.append(occurrence.nom)
    return noms


def _remplacer_dans_paragraphe(paragraphe, valeurs: Dict[str, str]) -> None:
    """
    Remplace les balises {{...}} d'un paragraphe.

    Cas normal (balise contenue dans un seul run, comme dans le modèle
    standard généré par templates/build_template.py) : remplacement run
    par run, la mise en forme de chaque run est conservée. Si une
    balise est scindée entre plusieurs runs (document retouché dans
    Word), le texte complet du paragraphe est reconstruit dans le
    premier run.
    """
    texte_complet = "".join(run.text for run in paragraphe.runs)
    if "{{" not in texte_complet:
        return

    for run in paragraphe.runs:
        if "{{" in run.text:
            run.text = remplacer_balises(run.text, valeurs)

    texte_apres = "".join(run.text for run in paragraphe.runs)
    table = valeurs_par_nom(valeurs)
    reste_remplacable = any(o.nom in table for o in trouver_balises(texte_apres))
    if not reste_remplacable:
        return

    # Filet de sécurité : balise scindée entre plusieurs runs
    texte_final = remplacer_balises(texte_apres, valeurs)
    if paragraphe.runs:
        paragraphe.runs[0].text = texte_final
        for run in paragraphe.runs[1:]:
            run.text = ""


def _remplacer_dans_document(document: Document, valeurs: Dict[str, str]) -> None:
    """Applique le remplacement des balises à tous les paragraphes du document (voir iterer_paragraphes)."""
    for paragraphe in iterer_paragraphes(document):
        _remplacer_dans_paragraphe(paragraphe, valeurs)


def verifier_aucun_placeholder_restant(document: Document) -> None:
    """
    Vérifie qu'aucune balise {{...}} ne subsiste dans le document.
    Lève WordExportError si une balise a été oubliée — un bulletin ne
    doit jamais être livré avec un champ non rempli.
    """
    for paragraphe in iterer_paragraphes(document):
        texte = paragraphe.text
        if contient_balise_residuelle(texte):
            raise WordExportError(f"Placeholder non remplacé détecté dans le bulletin : {texte!r}")


def generer_document_bulletin(valeurs: Dict[str, str], template_path: Optional[Path] = None) -> Document:
    """
    Charge le template officiel et retourne un objet Document avec
    tous les placeholders remplacés par `valeurs`. Ne sauvegarde rien
    sur disque (cf. `sauvegarder_document`).
    """
    chemin_template = template_path if template_path is not None else TEMPLATE_PATH
    if not chemin_template.exists():
        raise WordExportError(
            f"Le template officiel est introuvable : {chemin_template}. "
            "Il doit être généré une fois via templates/build_template.py."
        )

    document = Document(chemin_template)
    _remplacer_dans_document(document, valeurs)
    verifier_aucun_placeholder_restant(document)
    return document


def generer_nom_fichier_bulletin(nom: str, prenom: str, libelle_periode: str, extension: str = ".docx") -> str:
    """
    Construit le nom de fichier dynamique et sûr sous Windows d'un
    bulletin individuel, ex : Bulletin_KAMGANG_JEAN_PAUL_AOUT_2026.docx
    (ou .pdf lorsque le modèle actif est un modèle PDF).
    """
    nom_nettoye = nettoyer_nom_fichier(nom.upper())
    prenom_nettoye = nettoyer_nom_fichier(prenom.upper())
    periode_nettoyee = nettoyer_nom_fichier(libelle_periode.upper())
    return f"Bulletin_{nom_nettoye}_{prenom_nettoye}_{periode_nettoyee}{extension}"


def sauvegarder_document(document: Document, chemin: Path) -> Path:
    """Enregistre le document Word au chemin donné (le dossier parent est créé si besoin)."""
    chemin.parent.mkdir(parents=True, exist_ok=True)
    document.save(chemin)
    return chemin


# ---------------------------------------------------------------------
# Bulletin Word déjà rempli -> modèle à balises
# ---------------------------------------------------------------------

@dataclass
class SegmentDocx(Segment):
    """Segment de texte d'un document Word : paragraphe n° `paragraphe`, caractères [debut, fin[."""

    paragraphe: int = 0
    debut: int = 0
    fin: int = 0


_LARGEUR_COLONNE = 10_000  # coordonnée symbolique : colonne × 10 000 + position dans le texte


def ouvrir_document(contenu: bytes) -> Document:
    try:
        return Document(io.BytesIO(bytes(contenu)))
    except Exception as erreur:  # noqa: BLE001 — tout échec d'ouverture = fichier invalide
        raise WordExportError(f"Le fichier n'est pas un document Word (.docx) lisible ({erreur}).") from erreur


def _grouper_mots(texte: str) -> List[Tuple[int, int]]:
    """Découpe un texte en groupes (début, fin) : nombres isolés, suites de mots, « mois année »."""
    mots = [(m.start(), m.end()) for m in re.finditer(r"\S+", texte)]
    groupes: List[List[Tuple[int, int]]] = []
    for debut, fin in mots:
        if groupes:
            prec_debut, prec_fin = groupes[-1][-1]
            ecart = texte[prec_fin:debut]
            a, b = texte[prec_debut:prec_fin], texte[debut:fin]
            genre_a, genre_b = genre_texte(a), genre_texte(b)
            fusion = "\t" not in ecart and len(ecart) <= 2 and not a.endswith(":")
            if fusion and (genre_a == "nombre" or genre_b == "nombre"):
                if genre_a == "nombre" and genre_b == "nombre":
                    fusion = bool(re.fullmatch(r"\d{3}([.,]\d+)?", b))
                else:
                    fusion = genre_a != "nombre" and bool(mois_annee(f"{a} {b}"))
            if fusion:
                groupes[-1].append((debut, fin))
                continue
        groupes.append([(debut, fin)])
    return [(groupe[0][0], groupe[-1][1]) for groupe in groupes]


def segments_docx(contenu: bytes) -> List[SegmentDocx]:
    """Segments de texte d'un bulletin Word, avec repères de ligne et de colonne."""
    document = ouvrir_document(contenu)
    segments: List[SegmentDocx] = []
    for index, (paragraphe, rang, colonne) in enumerate(_parcourir(document)):
        texte = paragraphe.text
        for debut, fin in _grouper_mots(texte):
            segments.append(SegmentDocx(
                id=f"w{index}-{debut}", texte=texte[debut:fin],
                x0=colonne * _LARGEUR_COLONNE + debut, x1=colonne * _LARGEUR_COLONNE + fin,
                y0=float(rang), y1=rang + 0.5, page=0, paragraphe=index, debut=debut, fin=fin,
            ))
    return segments


def convertir_en_modele(contenu: bytes, correspondances: Dict[str, str]) -> bytes:
    """
    Transforme un bulletin Word rempli en modèle à balises : chaque
    segment retenu ({segment.id: "NOM_BALISE"}) est remplacé par
    {{NOM_BALISE}}. La mise en forme est conservée lorsque le texte
    remplacé tient dans un seul run (cas habituel).
    """
    segments = {s.id: s for s in segments_docx(contenu)}
    inconnus = [cle for cle in correspondances if cle not in segments]
    if inconnus:
        raise WordExportError("Le document a changé depuis l'analyse : relancez l'analyse du modèle.")

    par_paragraphe: Dict[int, List[Tuple[int, int, str]]] = {}
    for cle, balise in correspondances.items():
        segment = segments[cle]
        par_paragraphe.setdefault(segment.paragraphe, []).append((segment.debut, segment.fin, balise))

    document = ouvrir_document(contenu)
    for index, paragraphe in enumerate(iterer_paragraphes(document)):
        for debut, fin, balise in sorted(par_paragraphe.get(index, []), reverse=True):
            _remplacer_plage(paragraphe, debut, fin, "{{" + balise + "}}")

    tampon = io.BytesIO()
    document.save(tampon)
    return tampon.getvalue()


def _remplacer_plage(paragraphe, debut: int, fin: int, nouveau: str) -> None:
    position = 0
    for run in paragraphe.runs:
        longueur = len(run.text)
        if position <= debut and fin <= position + longueur:
            run.text = run.text[: debut - position] + nouveau + run.text[fin - position:]
            return
        position += longueur
    # Plage répartie sur plusieurs runs : texte reconstruit dans le premier run.
    texte = paragraphe.text
    if paragraphe.runs:
        paragraphe.runs[0].text = texte[:debut] + nouveau + texte[fin:]
        for run in paragraphe.runs[1:]:
            run.text = ""


def document_en_octets(document: Document) -> bytes:
    tampon = io.BytesIO()
    document.save(tampon)
    return tampon.getvalue()
