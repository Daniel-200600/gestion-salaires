"""
Service central d'importation massive de données (module 15).

Indépendant de Streamlit : utilisable depuis l'interface, un script,
ou une future API. Réutilise exclusivement les validations déjà
existantes (utils/validators.py) et les repositories déjà existants —
aucune deuxième source de vérité pour les règles métier ou les
formules de paie (services/paie_service.py reste l'unique moteur de
calcul, jamais dupliqué ici).

Pipeline (section 2) :

    Fichier -> Lecture -> Détection format -> Analyse colonnes ->
    Prévisualisation -> Validation technique -> Validation métier ->
    Détection doublons -> Dry run -> Confirmation -> Import
    transactionnel -> Rapport -> Audit

Chaque étape est une fonction pure ou quasi-pure, testable
indépendamment.
"""

import logging
import re
import unicodedata
from difflib import SequenceMatcher
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union

import pandas as pd

from database.connection import get_connection
from database.repositories import (
    enseignant_repository,
    heures_repository,
    import_repository,
    periode_repository,
)
from models.audit_log import AuditLog
from models.enseignant import Enseignant
from models.enums import (
    NiveauAnomalie,
    StatutImport,
    StatutPeriode,
    StrategieDoublon,
    TypeActionAudit,
    TypeImport,
)
from models.import_journal import AnomalieImport, ImportJournal
from models.saisie_heures import SaisieHeures
from utils.validators import (
    EnseignantValidationError,
    message_fiche_incomplete,
    nettoyer_champ_optionnel,
    nettoyer_texte,
    valider_heures,
    valider_nom_ou_prenom,
    valider_sexe,
    valider_statut,
    valider_taux_horaire,
)

DbPath = Optional[Union[str, Path]]

logger = logging.getLogger("salaires_app.import_service")

EXTENSIONS_AUTORISEES = {".xlsx", ".csv", ".docx", ".pdf"}
FORMATS_ACCEPTES = "Excel (.xlsx), CSV (.csv), Word (.docx) ou PDF (.pdf)"
TAILLE_MAX_OCTETS = 10 * 1024 * 1024  # 10 Mo — documenté (section 5)


class ImportServiceError(Exception):
    """Erreur de fichier ou de configuration d'import — message toujours compréhensible."""


# ---------------------------------------------------------------------
# Étape A/B — Chargement et lecture (section 5/6)
# ---------------------------------------------------------------------

@dataclass
class ContenuFichier:
    """Résultat de la lecture d'un fichier (section 6) : une ou plusieurs feuilles."""

    nom_fichier: str
    feuilles: Dict[str, pd.DataFrame] = field(default_factory=dict)

    @property
    def noms_feuilles(self) -> List[str]:
        return list(self.feuilles.keys())


def lire_fichier(chemin: Path, nom_fichier: Optional[str] = None) -> ContenuFichier:
    """
    Lit un fichier .xlsx (toutes feuilles), .csv (une seule « feuille »
    nommée `csv`), .docx (chaque tableau du document) ou .pdf (chaque
    tableau, y compris une liste qui se poursuit sur plusieurs pages). Contrôle systématiquement : extension,
    existence, taille, lisibilité, présence de données (section 5).

    Ne fait jamais confiance au nom de fichier fourni par
    l'utilisateur pour construire un chemin : `chemin` doit toujours
    être un fichier déjà présent sur un emplacement contrôlé par
    l'application (ex. un fichier temporaire créé par Streamlit),
    jamais un chemin arbitraire assemblé à partir d'une saisie libre.
    """
    nom_affiche = nom_fichier or chemin.name
    extension = Path(nom_affiche).suffix.lower()

    if extension not in EXTENSIONS_AUTORISEES:
        raise ImportServiceError(
            f"Extension non autorisée : « {extension} ». Formats acceptés : {FORMATS_ACCEPTES}."
        )
    if not chemin.exists():
        raise ImportServiceError("Le fichier est introuvable.")
    taille = chemin.stat().st_size
    if taille == 0:
        raise ImportServiceError("Le fichier est vide.")
    if taille > TAILLE_MAX_OCTETS:
        raise ImportServiceError(
            f"Le fichier dépasse la taille maximale autorisée ({TAILLE_MAX_OCTETS // (1024*1024)} Mo)."
        )

    try:
        if extension == ".xlsx":
            classeur = pd.ExcelFile(chemin)
            if not classeur.sheet_names:
                raise ImportServiceError("Le classeur Excel ne contient aucune feuille.")
            feuilles = {}
            for nom_feuille in classeur.sheet_names:
                df = classeur.parse(nom_feuille, dtype=str)
                feuilles[nom_feuille] = df
            if all(df.empty for df in feuilles.values()):
                raise ImportServiceError(
                    "Aucune donnée exploitable dans le classeur (toutes les feuilles sont vides)."
                )
        elif extension == ".docx":
            feuilles = _lire_tableaux_word(chemin)
        elif extension == ".pdf":
            feuilles = _lire_tableaux_pdf(chemin)
        else:  # .csv
            df = pd.read_csv(chemin, dtype=str)
            if df.empty:
                raise ImportServiceError("Le fichier CSV ne contient aucune ligne de données.")
            feuilles = {"csv": df}
    except ImportServiceError:
        raise
    except (pd.errors.ParserError, pd.errors.EmptyDataError, ValueError, UnicodeDecodeError) as erreur:
        raise ImportServiceError(f"Fichier illisible ou corrompu : {erreur}") from erreur
    except Exception as erreur:  # noqa: BLE001 — jamais de traceback brut à l'utilisateur
        logger.error("Erreur inattendue à la lecture de %s : %s", nom_affiche, erreur)
        raise ImportServiceError("Le fichier n'a pas pu être lu (format inattendu).") from erreur

    return ContenuFichier(nom_fichier=nom_affiche, feuilles=feuilles)


# ---------------------------------------------------------------------
# Lecture des listes Word et PDF : chaque tableau devient une « feuille »
# ---------------------------------------------------------------------

def _tableau_vers_dataframe(lignes: List[List[Optional[str]]]) -> pd.DataFrame:
    """
    Première ligne non vide = en-têtes ; lignes entièrement vides ignorées.
    Un en-tête vide ou répété reçoit un nom distinct (« Colonne 3 »...).
    """
    lignes = [[nettoyer_texte(str(c)) if c is not None else "" for c in ligne] for ligne in lignes]
    lignes = [ligne for ligne in lignes if any(ligne)]
    if len(lignes) < 2:
        return pd.DataFrame()
    largeur = max(len(ligne) for ligne in lignes)
    lignes = [ligne + [""] * (largeur - len(ligne)) for ligne in lignes]
    en_tetes, vus = [], set()
    for position, en_tete in enumerate(lignes[0], start=1):
        nom = en_tete or f"Colonne {position}"
        while nom in vus:
            nom = f"{nom} ({position})"
        vus.add(nom)
        en_tetes.append(nom)
    return pd.DataFrame(lignes[1:], columns=en_tetes, dtype=str).replace("", None)


def _lire_tableaux_word(chemin: Path) -> Dict[str, pd.DataFrame]:
    from docx import Document

    try:
        document = Document(str(chemin))
    except Exception as erreur:  # noqa: BLE001
        raise ImportServiceError("Le document Word n'a pas pu être ouvert (fichier endommagé ?).") from erreur
    feuilles = {}
    for numero, tableau in enumerate(document.tables, start=1):
        df = _tableau_vers_dataframe([[cellule.text for cellule in ligne.cells] for ligne in tableau.rows])
        if not df.empty:
            feuilles[f"Tableau {numero}"] = df
    if not feuilles:
        raise ImportServiceError(
            "Aucun tableau exploitable dans le document Word : présentez la liste sous forme de tableau "
            "(une ligne par enseignant, les intitulés des colonnes sur la première ligne)."
        )
    return feuilles


def _ressemble_a_un_en_tete(ligne: List[Optional[str]]) -> bool:
    champs = {_cle_normalisation(v) for variantes in VARIANTES_COLONNES.values() for vs in variantes.values() for v in vs}
    return any(_cle_normalisation(c or "") in champs for c in ligne)


def _lire_tableaux_pdf(chemin: Path) -> Dict[str, pd.DataFrame]:
    import pymupdf

    try:
        document = pymupdf.open(str(chemin))
    except Exception as erreur:  # noqa: BLE001
        raise ImportServiceError("Le PDF n'a pas pu être ouvert (fichier endommagé ou protégé ?).") from erreur
    if not any(page.get_text().strip() for page in document):
        raise ImportServiceError(
            "Ce PDF ne contient pas de texte (document scanné ou photographié) : il ne peut pas être lu. "
            "Utilisez le fichier Word ou Excel d'origine."
        )
    tableaux: List[List[List[Optional[str]]]] = []
    for page in document:
        for tableau in page.find_tables().tables:
            lignes = tableau.extract()
            if not lignes:
                continue
            precedent = tableaux[-1] if tableaux else None
            # Une liste longue continue sur la page suivante : même nombre de
            # colonnes, sans en-tête (ou avec le même en-tête répété).
            if precedent is not None and len(lignes[0]) == len(precedent[0]) and (
                lignes[0] == precedent[0] or not _ressemble_a_un_en_tete(lignes[0])
            ):
                precedent.extend(lignes[1:] if lignes[0] == precedent[0] else lignes)
            else:
                tableaux.append([list(ligne) for ligne in lignes])
    feuilles = {}
    for numero, lignes in enumerate(tableaux, start=1):
        df = _tableau_vers_dataframe(lignes)
        if not df.empty:
            feuilles[f"Tableau {numero}"] = df
    if not feuilles:
        raise ImportServiceError(
            "Aucun tableau reconnu dans le PDF : la liste doit être présentée sous forme de tableau "
            "(une ligne par enseignant)."
        )
    return feuilles


# ---------------------------------------------------------------------
# Normalisation des colonnes (section 7)
# ---------------------------------------------------------------------

VARIANTES_COLONNES: Dict[TypeImport, Dict[str, List[str]]] = {
    TypeImport.ENSEIGNANTS: {
        "nom": ["nom", "noms", "nom enseignant", "nom_enseignant", "nom de famille", "last name"],
        "prenom": ["prenom", "prénom", "prenoms", "prénoms", "prenom s", "prenom enseignant", "prenom_enseignant",
                   "first name"],
        # Une seule colonne pour le nom et le(s) prénom(s) : séparée à l'import.
        "nom_complet": ["nom et prenom", "nom et prenoms", "noms et prenoms", "nom prenom", "nom prenoms",
                        "noms prenoms", "nom complet", "nom et prenom s", "noms et prenom s", "enseignant",
                        "enseignants", "identite", "nom de l enseignant", "nom et prenoms de l enseignant"],
        "sexe": ["sexe", "genre", "sexe m f", "m f", "h f"],
        "statut": ["statut", "statut enseignant", "statut v p", "v p", "categorie", "type", "situation",
                   "vacataire permanent", "type de contrat", "contrat"],
        "taux_horaire": ["taux_horaire", "taux horaire", "taux", "taux horaire fcfa", "taux fcfa", "taux heure",
                         "taux par heure", "taux h", "prix de l heure", "prix horaire", "montant horaire"],
        "email": ["email", "e-mail", "courriel", "mail", "adresse email", "adresse e-mail", "adresse mail"],
        "telephone": ["telephone", "téléphone", "tel", "tél", "contact", "contacts", "numero", "numero de telephone",
                      "n telephone", "n tel", "portable", "mobile", "cellulaire"],
        "adresse": ["adresse", "domicile", "lieu de residence", "residence", "quartier", "ville"],
    },
    TypeImport.HEURES: {
        "nom": ["nom", "nom enseignant", "nom_enseignant"],
        "prenom": ["prenom", "prénom", "prenom enseignant", "prenom_enseignant"],
        "semaine_1": ["semaine_1", "semaine 1", "s1"],
        "semaine_2": ["semaine_2", "semaine 2", "s2"],
        "semaine_3": ["semaine_3", "semaine 3", "s3"],
        "semaine_4": ["semaine_4", "semaine 4", "s4"],
        "semaine_5": ["semaine_5", "semaine 5", "s5"],
    },
    TypeImport.REMUNERATIONS: {
        "nom": ["nom", "nom enseignant", "nom_enseignant"],
        "prenom": ["prenom", "prénom", "prenom enseignant", "prenom_enseignant"],
        "prime_ap_pp": ["prime_ap_pp", "ap/pp", "ap_pp"],
        "surveillance_secretariat": ["surveillance_secretariat", "surveillance/secretariat", "surveillance"],
        "indemnite_suggestion_admin": ["indemnite_suggestion_admin", "indemnite", "indemnité"],
    },
    TypeImport.RETENUES: {
        "nom": ["nom", "nom enseignant", "nom_enseignant"],
        "prenom": ["prenom", "prénom", "prenom enseignant", "prenom_enseignant"],
        "retenue_amicale": ["retenue_amicale", "retenue amicale"],
        "dette": ["dette"],
    },
}


def _cle_normalisation(valeur: str) -> str:
    """
    Forme de comparaison d'un intitulé ou d'un nom : sans accents, en
    minuscules, ponctuation remplacée par des espaces (« Prénom(s) » ->
    « prenom s », « Taux/heure » -> « taux heure »).
    """
    texte = unicodedata.normalize("NFKD", str(valeur))
    texte = "".join(c for c in texte if not unicodedata.combining(c)).lower()
    return " ".join(re.sub(r"[^0-9a-z]+", " ", texte).split())


@dataclass
class ResultatNormalisation:
    correspondance: Dict[str, str]  # colonne brute -> champ canonique
    colonnes_inconnues: List[str]


def normaliser_colonnes(colonnes_brutes: List[str], type_import: TypeImport) -> ResultatNormalisation:
    """
    Associe chaque colonne du fichier à un champ métier canonique
    (section 7). Une colonne inconnue n'est JAMAIS associée par
    approximation dangereuse : elle est listée telle quelle parmi les
    colonnes non reconnues, visible par l'utilisateur.
    """
    variantes = VARIANTES_COLONNES[type_import]
    lookup = {}
    for champ, variantes_champ in variantes.items():
        for variante in variantes_champ:
            lookup[_cle_normalisation(variante)] = champ

    correspondance: Dict[str, str] = {}
    colonnes_inconnues: List[str] = []
    for colonne in colonnes_brutes:
        cle = _cle_normalisation(colonne)
        # « Taux horaire (FCFA) », « Taux/heure en F CFA » : la devise en fin
        # d'intitulé ne change pas la nature de la colonne.
        sans_devise = re.sub(r"( en)?( f cfa| fcfa| xaf| cfa| f)+$", "", cle)
        if cle in lookup:
            correspondance[colonne] = lookup[cle]
        elif sans_devise in lookup:
            correspondance[colonne] = lookup[sans_devise]
        else:
            colonnes_inconnues.append(colonne)

    return ResultatNormalisation(correspondance=correspondance, colonnes_inconnues=colonnes_inconnues)


# ---------------------------------------------------------------------
# Interprétation souple des valeurs d'une liste existante (import enseignants)
# ---------------------------------------------------------------------

_SEXES = {"m": "M", "masculin": "M", "homme": "M", "h": "M", "garcon": "M", "male": "M",
          "f": "F", "feminin": "F", "femme": "F", "fille": "F", "female": "F"}
_STATUTS = {"v": "V", "vacataire": "V", "vac": "V", "vacation": "V",
            "p": "P", "permanent": "P", "perm": "P", "titulaire": "P", "permanente": "P"}


def interpreter_sexe(valeur: Optional[str]):
    """« Masculin », « H », « F », « Femme »... -> Sexe, ou None si non reconnu."""
    code = _SEXES.get(_cle_normalisation(valeur or ""))
    return valider_sexe(code) if code else None


def interpreter_statut(valeur: Optional[str]):
    """« Vacataire », « V », « Permanent »... -> StatutEnseignant, ou None si non reconnu."""
    code = _STATUTS.get(_cle_normalisation(valeur or ""))
    return valider_statut(code) if code else None


def interpreter_taux_horaire(valeur: Optional[str]) -> Optional[int]:
    """« 1 500 », « 1500 FCFA », « 1.500 », « 1500,00 » -> 1500 ; None si non reconnu."""
    if valeur is None:
        return None
    texte = re.sub(r"(?i)\s*(f\s*cfa|xaf|fcfa|francs?|f)\s*$", "", str(valeur).strip())
    texte = re.sub(r"[\s\u00a0\u202f]", "", texte)
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", texte):
        texte = texte.replace(".", "")
    texte = re.sub(r"[,.]0+$", "", texte)
    try:
        return valider_taux_horaire(texte)
    except EnseignantValidationError:
        return None


def separer_nom_complet(valeur: str) -> Tuple[str, str]:
    """
    « MBARGA Élise » -> ("MBARGA", "Élise") ; « MBALLA NDOME Clarisse » ->
    ("MBALLA NDOME", "Clarisse"). Les mots en majuscules du début forment le
    nom ; à défaut (tout en majuscules, ou aucune majuscule), le premier mot.
    """
    mots = nettoyer_texte(valeur).split()
    if not mots:
        return "", ""
    debut = 0
    while debut < len(mots) and mots[debut].isupper():
        debut += 1
    if 0 < debut < len(mots):
        return " ".join(mots[:debut]), " ".join(mots[debut:])
    return mots[0], " ".join(mots[1:])


# ---------------------------------------------------------------------
# Étape analyse / prévisualisation (section 10)
# ---------------------------------------------------------------------

@dataclass
class RapportAnalyse:
    nom_fichier: str
    feuille_choisie: str
    noms_feuilles: List[str]
    nombre_lignes: int
    colonnes_detectees: List[str]
    colonnes_reconnues: Dict[str, str]
    colonnes_inconnues: List[str]
    apercu: List[dict] = field(default_factory=list)


def analyser_fichier(
    contenu: ContenuFichier, type_import: TypeImport, feuille: Optional[str] = None
) -> RapportAnalyse:
    """Analyse une feuille du fichier chargé (section 6/10) — ne modifie jamais la base."""
    nom_feuille = feuille or contenu.noms_feuilles[0]
    if nom_feuille not in contenu.feuilles:
        raise ImportServiceError(f"Feuille introuvable : « {nom_feuille} ».")

    df = contenu.feuilles[nom_feuille]
    if df.empty:
        raise ImportServiceError(f"La feuille « {nom_feuille} » ne contient aucune donnée exploitable.")

    colonnes = [str(c) for c in df.columns]
    normalisation = normaliser_colonnes(colonnes, type_import)

    return RapportAnalyse(
        nom_fichier=contenu.nom_fichier,
        feuille_choisie=nom_feuille,
        noms_feuilles=contenu.noms_feuilles,
        nombre_lignes=len(df),
        colonnes_detectees=colonnes,
        colonnes_reconnues=normalisation.correspondance,
        colonnes_inconnues=normalisation.colonnes_inconnues,
        apercu=df.head(10).fillna("").to_dict(orient="records"),
    )


# ---------------------------------------------------------------------
# Ligne d'import — structure commune à tous les types
# ---------------------------------------------------------------------

class ActionLigne:
    CREATION = "creation"
    MISE_A_JOUR = "mise_a_jour"
    IGNOREE = "ignoree"
    REJETEE = "rejetee"


@dataclass
class LigneImport:
    numero_ligne: int  # 1-based, correspond à la ligne du fichier (hors en-tête)
    donnees: dict
    anomalies: List[AnomalieImport] = field(default_factory=list)
    action: str = ActionLigne.REJETEE
    cible_id: Optional[int] = None  # id de l'enregistrement existant en cas de mise à jour
    doublon_probable: Optional[str] = None  # enseignant au nom très proche (base ou fichier)

    @property
    def a_une_erreur(self) -> bool:
        return any(a.niveau == NiveauAnomalie.ERREUR for a in self.anomalies)


def _ajouter_erreur(ligne: LigneImport, champ: Optional[str], valeur, message: str) -> None:
    ligne.anomalies.append(AnomalieImport(
        ligne=ligne.numero_ligne, niveau=NiveauAnomalie.ERREUR, message=message, champ=champ,
        valeur=str(valeur) if valeur is not None else None,
    ))


def _ajouter_avertissement(ligne: LigneImport, champ: Optional[str], valeur, message: str) -> None:
    ligne.anomalies.append(AnomalieImport(
        ligne=ligne.numero_ligne, niveau=NiveauAnomalie.AVERTISSEMENT, message=message, champ=champ,
        valeur=str(valeur) if valeur is not None else None,
    ))


def _valeur_normalisee(row: dict, correspondance: Dict[str, str], champ: str):
    for colonne_brute, champ_canonique in correspondance.items():
        if champ_canonique == champ:
            valeur = row.get(colonne_brute)
            if valeur is None or (isinstance(valeur, float) and pd.isna(valeur)):
                return None
            valeur_str = str(valeur).strip()
            return valeur_str if valeur_str else None
    return None


@dataclass
class RapportPreparation:
    """Résultat de la préparation (validation + doublons), commun au dry-run et à l'import réel."""

    lignes: List[LigneImport] = field(default_factory=list)

    @property
    def nb_lignes(self) -> int:
        return len(self.lignes)

    @property
    def nb_creations(self) -> int:
        return sum(1 for l in self.lignes if l.action == ActionLigne.CREATION)

    @property
    def nb_mises_a_jour(self) -> int:
        return sum(1 for l in self.lignes if l.action == ActionLigne.MISE_A_JOUR)

    @property
    def nb_ignorees(self) -> int:
        return sum(1 for l in self.lignes if l.action == ActionLigne.IGNOREE)

    @property
    def nb_rejetees(self) -> int:
        return sum(1 for l in self.lignes if l.action == ActionLigne.REJETEE)

    @property
    def nb_erreurs(self) -> int:
        return sum(1 for l in self.lignes for a in l.anomalies if a.niveau == NiveauAnomalie.ERREUR)

    @property
    def nb_avertissements(self) -> int:
        return sum(1 for l in self.lignes for a in l.anomalies if a.niveau == NiveauAnomalie.AVERTISSEMENT)

    @property
    def nb_fiches_incompletes(self) -> int:
        """Enseignants créés ou mis à jour dont la fiche restera à compléter."""
        return sum(
            1 for l in self.lignes
            if l.action in (ActionLigne.CREATION, ActionLigne.MISE_A_JOUR) and l.donnees.get("a_completer")
        )

    @property
    def nb_doublons_probables(self) -> int:
        return sum(1 for l in self.lignes if l.doublon_probable)

    @property
    def toutes_anomalies(self) -> List[AnomalieImport]:
        return [a for l in self.lignes for a in l.anomalies]


# ---------------------------------------------------------------------
# Préparation — Import ENSEIGNANTS
# ---------------------------------------------------------------------

SEUIL_SIMILARITE_NOMS = 0.88


def _cle_identite(nom: str, prenom: str) -> tuple:
    """
    Identité d'un enseignant pour la détection des doublons : les mots du
    nom et du prénom, sans accents ni casse, dans n'importe quel ordre
    (« MBARGA Élise » = « Elise Mbarga » = nom complet « Mbarga Elise »).
    """
    return tuple(sorted(_cle_normalisation(f"{nom} {prenom}").split()))


def _sont_probablement_la_meme_personne(cle_a: tuple, cle_b: tuple) -> bool:
    """Un prénom en plus ou en moins (« NGONO Marie » / « NGONO Marie Claire ») ou une faute de frappe."""
    petite, grande = sorted((set(cle_a), set(cle_b)), key=len)
    if len(petite) >= 2 and petite < grande:
        return True
    return SequenceMatcher(None, " ".join(cle_a), " ".join(cle_b)).ratio() >= SEUIL_SIMILARITE_NOMS

def preparer_import_enseignants(
    analyse: RapportAnalyse,
    df: pd.DataFrame,
    strategie_doublon: StrategieDoublon = StrategieDoublon.REFUSER,
    db_path: DbPath = None,
    ignorer_doublons_probables: bool = False,
) -> RapportPreparation:
    """
    Valide chaque ligne et détecte les doublons (internes au fichier et
    contre la base) — SANS jamais écrire en base (section 13/14). Une seule
    requête pour tous les enseignants existants (section 26).

    Un même nom écrit dans un autre ordre est reconnu comme le même
    enseignant. Un nom très proche sans être identique (prénom en plus ou
    en moins, faute de frappe) est signalé comme doublon probable : la
    ligne est créée avec un avertissement, ou écartée si
    `ignorer_doublons_probables` est vrai.

    Seul le nom est obligatoire. Un sexe, un statut ou un taux horaire
    absent ou illisible n'écarte pas la ligne : l'enseignant est créé avec
    une fiche « à compléter » (avertissement), exclue de la paie jusqu'à ce
    qu'elle soit complétée dans Gestion › Enseignants. Avec la stratégie
    METTRE_A_JOUR, seules les informations présentes dans le fichier sont
    reprises : une valeur déjà connue n'est jamais effacée.
    """
    enseignants_existants = enseignant_repository.lister(inclure_inactifs=True, db_path=db_path)
    index_existants = {_cle_identite(e.nom, e.prenom): e for e in enseignants_existants}
    # Noms déjà connus, pour la recherche des doublons probables : base puis lignes du fichier.
    noms_connus: List[tuple] = [
        (cle, f"« {e.nom} {e.prenom}".strip() + f" » (id {e.id}, déjà enregistré)")
        for cle, e in index_existants.items()
    ]
    colonnes = analyse.colonnes_reconnues

    rapport = RapportPreparation()
    cles_vues_dans_fichier: Dict[tuple, int] = {}

    for position, row in enumerate(df.to_dict(orient="records"), start=1):
        ligne = LigneImport(numero_ligne=position, donnees=row)
        valeur = lambda champ: _valeur_normalisee(row, colonnes, champ)  # noqa: E731

        nom_brut, prenom_brut = valeur("nom"), valeur("prenom")
        if nom_brut is None and valeur("nom_complet") is not None:
            nom_brut, prenom_separe = separer_nom_complet(valeur("nom_complet"))
            prenom_brut = prenom_brut or prenom_separe
        nom = nettoyer_texte(nom_brut)
        prenom = nettoyer_texte(prenom_brut)
        if not nom:
            if any(v is not None for v in row.values() if not (isinstance(v, float) and pd.isna(v))):
                _ajouter_erreur(ligne, "nom", nom_brut, "Nom absent : la ligne ne peut pas être importée.")
                rapport.lignes.append(ligne)
            continue

        donnees = {"nom": nom, "prenom": prenom}
        for champ, interpreter, libelle in (
            ("sexe", interpreter_sexe, "sexe"),
            ("statut", interpreter_statut, "statut"),
            ("taux_horaire", interpreter_taux_horaire, "taux horaire"),
        ):
            brut = valeur(champ)
            donnees[champ] = interpreter(brut)
            if brut is not None and donnees[champ] is None:
                _ajouter_avertissement(ligne, champ, brut, f"{libelle.capitalize()} « {brut} » non reconnu : à compléter.")
        for champ in ("email", "telephone", "adresse"):
            donnees[champ] = nettoyer_champ_optionnel(valeur(champ))

        cle = _cle_identite(nom, prenom)
        nom_affiche = f"{nom} {prenom}".strip()
        if cle in cles_vues_dans_fichier:
            _ajouter_avertissement(
                ligne, "nom", nom_affiche,
                f"Doublon interne au fichier (déjà présent à la ligne {cles_vues_dans_fichier[cle]}).",
            )
            ligne.action = ActionLigne.IGNOREE
            rapport.lignes.append(ligne)
            continue
        cles_vues_dans_fichier[cle] = position

        existant = index_existants.get(cle)
        if existant is None:
            proche = next((libelle for cle_connue, libelle in noms_connus
                           if _sont_probablement_la_meme_personne(cle, cle_connue)), None)
            noms_connus.append((cle, f"« {nom_affiche} » (ligne {position} du fichier)"))
            if proche is not None:
                ligne.doublon_probable = proche
                if ignorer_doublons_probables:
                    _ajouter_avertissement(
                        ligne, "nom", nom_affiche, f"Doublon probable de {proche} — ligne écartée."
                    )
                    ligne.action = ActionLigne.IGNOREE
                    rapport.lignes.append(ligne)
                    continue
                _ajouter_avertissement(
                    ligne, "nom", nom_affiche,
                    f"Doublon probable de {proche} : vérifiez qu'il ne s'agit pas du même enseignant.",
                )
        if existant is not None:
            if strategie_doublon == StrategieDoublon.REFUSER:
                _ajouter_erreur(
                    ligne, "nom", nom_affiche,
                    f"« {nom_affiche} » existe déjà en base (id {existant.id}) — stratégie REFUSER.",
                )
                rapport.lignes.append(ligne)
                continue
            if strategie_doublon == StrategieDoublon.IGNORER:
                _ajouter_avertissement(ligne, "nom", nom_affiche, "Enseignant déjà existant — ignoré (stratégie IGNORER).")
                ligne.action = ActionLigne.IGNOREE
                rapport.lignes.append(ligne)
                continue
            # METTRE_A_JOUR : complète sans jamais effacer une valeur connue.
            conflits = [champ for champ in ("sexe", "statut", "taux_horaire")
                        if donnees[champ] is not None and getattr(existant, champ) not in (None, donnees[champ])]
            if conflits:
                _ajouter_avertissement(
                    ligne, "nom", nom_affiche,
                    "Valeurs différentes de la base (" + ", ".join(conflits) + ") — seront mises à jour.",
                )
            for champ in ("sexe", "statut", "taux_horaire", "email", "telephone", "adresse"):
                if donnees[champ] is None:
                    donnees[champ] = getattr(existant, champ)
            # Même personne, nom écrit autrement (ordre, accents) : la fiche garde son nom.
            donnees["nom"], donnees["prenom"] = existant.nom, existant.prenom
            ligne.action = ActionLigne.MISE_A_JOUR
            ligne.cible_id = existant.id
        else:
            ligne.action = ActionLigne.CREATION

        manquants = [libelle for champ, libelle in (("sexe", "sexe"), ("statut", "statut"),
                                                    ("taux_horaire", "taux horaire")) if donnees[champ] is None]
        donnees["a_completer"] = bool(manquants)
        if manquants:
            _ajouter_avertissement(
                ligne, None, nom_affiche,
                "Fiche à compléter (" + ", ".join(manquants) + ") : exclue de la paie tant qu'elle est incomplète.",
            )
        ligne.donnees = donnees
        rapport.lignes.append(ligne)

    return rapport


def executer_import_enseignants(
    rapport_preparation: RapportPreparation,
    nom_fichier: str,
    utilisateur: Optional[str] = None,
    db_path: DbPath = None,
) -> ImportJournal:
    """
    Import réel, transactionnel (section 16) : une seule connexion
    pour toutes les lignes ; ROLLBACK complet en cas d'erreur
    inattendue, COMMIT uniquement si tout s'est déroulé normalement.
    Les lignes REJETEE (déjà écartées lors de la préparation) ne sont
    jamais insérées.
    """
    journal = ImportJournal(
        nom_fichier=nom_fichier, type_import=TypeImport.ENSEIGNANTS, statut=StatutImport.ECHEC,
        utilisateur=utilisateur, nb_lignes=rapport_preparation.nb_lignes,
    )

    with get_connection(db_path) as conn:
        try:
            for ligne in rapport_preparation.lignes:
                if ligne.action not in (ActionLigne.CREATION, ActionLigne.MISE_A_JOUR):
                    continue
                d = ligne.donnees
                enseignant = Enseignant(
                    id=ligne.cible_id, nom=d["nom"], prenom=d["prenom"], sexe=d["sexe"], statut=d["statut"],
                    taux_horaire=d["taux_horaire"], email=d.get("email"), telephone=d.get("telephone"),
                    adresse=d.get("adresse"), actif=True,
                )
                if ligne.action == ActionLigne.CREATION:
                    enseignant_repository.creer(enseignant, conn=conn)
                else:
                    enseignant_repository.mettre_a_jour(enseignant, conn=conn)
        except Exception as erreur:  # noqa: BLE001
            conn.rollback()
            logger.error("Import enseignants annulé (ROLLBACK) : %s", erreur)
            journal.statut = StatutImport.ECHEC
            journal.nb_erreurs = rapport_preparation.nb_erreurs + 1
            journal.id = import_repository.enregistrer_import(journal, db_path=db_path)
            import_repository.enregistrer_anomalies(journal.id, rapport_preparation.toutes_anomalies, db_path=db_path)
            _journaliser_audit_import(journal, TypeActionAudit.IMPORT_DONNEES_ECHEC, db_path=db_path)
            return journal
        else:
            conn.commit()

    journal.statut = StatutImport.TERMINE
    journal.nb_creations = rapport_preparation.nb_creations
    journal.nb_mises_a_jour = rapport_preparation.nb_mises_a_jour
    journal.nb_ignorees = rapport_preparation.nb_ignorees
    journal.nb_rejetees = rapport_preparation.nb_rejetees
    journal.nb_erreurs = rapport_preparation.nb_erreurs
    journal.nb_avertissements = rapport_preparation.nb_avertissements
    journal.id = import_repository.enregistrer_import(journal, db_path=db_path)
    import_repository.enregistrer_anomalies(journal.id, rapport_preparation.toutes_anomalies, db_path=db_path)
    _journaliser_audit_import(journal, TypeActionAudit.IMPORT_DONNEES, db_path=db_path)
    return journal


# ---------------------------------------------------------------------
# Préparation — Import HEURES (respecte le verrouillage des périodes, module 12)
# ---------------------------------------------------------------------

def preparer_import_heures(
    analyse: RapportAnalyse, df: pd.DataFrame, periode_id: int, db_path: DbPath = None
) -> RapportPreparation:
    """
    Valide chaque ligne d'heures. Refuse (erreur par ligne concernée)
    si la période n'autorise pas la modification (module 12) — ne
    contourne jamais cette protection.
    """
    periode = periode_repository.obtenir_par_id(periode_id, db_path=db_path)
    if periode is None:
        raise ImportServiceError(f"Aucune période avec l'id {periode_id}.")
    periode_modifiable = periode.statut == StatutPeriode.OUVERTE

    enseignants_existants = enseignant_repository.lister(inclure_inactifs=True, db_path=db_path)
    index_existants = {(_cle_normalisation(e.nom), _cle_normalisation(e.prenom)): e for e in enseignants_existants}

    rapport = RapportPreparation()
    for position, row in enumerate(df.to_dict(orient="records"), start=1):
        ligne = LigneImport(numero_ligne=position, donnees=row)

        nom_brut = _valeur_normalisee(row, analyse.colonnes_reconnues, "nom")
        prenom_brut = _valeur_normalisee(row, analyse.colonnes_reconnues, "prenom")
        try:
            nom = valider_nom_ou_prenom(nom_brut, "nom")
            prenom = valider_nom_ou_prenom(prenom_brut, "prenom")
        except EnseignantValidationError as erreur:
            _ajouter_erreur(ligne, "nom", nom_brut, str(erreur))
            rapport.lignes.append(ligne)
            continue

        cle = (_cle_normalisation(nom), _cle_normalisation(prenom))
        enseignant = index_existants.get(cle)
        if enseignant is None:
            _ajouter_erreur(ligne, "nom", f"{nom} {prenom}", f"Enseignant inconnu : « {nom} {prenom} ».")
            rapport.lignes.append(ligne)
            continue
        if not enseignant.est_complet:
            _ajouter_erreur(ligne, "nom", f"{nom} {prenom}", message_fiche_incomplete(enseignant))
            rapport.lignes.append(ligne)
            continue

        if not periode_modifiable:
            _ajouter_erreur(
                ligne, "periode", periode.libelle,
                f"Cette période est {periode.statut.value.upper()}. Les données ne peuvent plus être modifiées.",
            )
            rapport.lignes.append(ligne)
            continue

        heures_par_semaine = {}
        erreur_heures = False
        for semaine in range(1, 6):
            champ = f"semaine_{semaine}"
            valeur_brute = _valeur_normalisee(row, analyse.colonnes_reconnues, champ)
            if valeur_brute is None:
                continue
            try:
                heures_par_semaine[semaine] = valider_heures(valeur_brute, champ)
            except ValueError as erreur:
                _ajouter_erreur(ligne, champ, valeur_brute, str(erreur))
                erreur_heures = True
        if erreur_heures:
            rapport.lignes.append(ligne)
            continue

        if not heures_par_semaine:
            _ajouter_avertissement(ligne, None, None, "Aucune heure renseignée sur cette ligne — ignorée.")
            ligne.action = ActionLigne.IGNOREE
            rapport.lignes.append(ligne)
            continue

        ligne.action = ActionLigne.MISE_A_JOUR  # upsert : création ou mise à jour selon l'existant
        ligne.cible_id = enseignant.id
        ligne.donnees = {"enseignant_id": enseignant.id, "heures_par_semaine": heures_par_semaine}
        rapport.lignes.append(ligne)

    return rapport


def executer_import_heures(
    rapport_preparation: RapportPreparation,
    nom_fichier: str,
    periode_id: int,
    utilisateur: Optional[str] = None,
    db_path: DbPath = None,
) -> ImportJournal:
    """Import transactionnel des heures — mêmes garanties que executer_import_enseignants."""
    journal = ImportJournal(
        nom_fichier=nom_fichier, type_import=TypeImport.HEURES, statut=StatutImport.ECHEC,
        utilisateur=utilisateur, periode_id=periode_id, nb_lignes=rapport_preparation.nb_lignes,
    )

    with get_connection(db_path) as conn:
        try:
            for ligne in rapport_preparation.lignes:
                if ligne.action != ActionLigne.MISE_A_JOUR:
                    continue
                for semaine, heures in ligne.donnees["heures_par_semaine"].items():
                    saisie = SaisieHeures(
                        enseignant_id=ligne.donnees["enseignant_id"], periode_id=periode_id,
                        numero_semaine=semaine, heures_effectuees=heures,
                    )
                    heures_repository.upsert(saisie, conn=conn)
        except Exception as erreur:  # noqa: BLE001
            conn.rollback()
            logger.error("Import heures annulé (ROLLBACK) : %s", erreur)
            journal.nb_erreurs = rapport_preparation.nb_erreurs + 1
            journal.id = import_repository.enregistrer_import(journal, db_path=db_path)
            import_repository.enregistrer_anomalies(journal.id, rapport_preparation.toutes_anomalies, db_path=db_path)
            _journaliser_audit_import(journal, TypeActionAudit.IMPORT_DONNEES_ECHEC, db_path=db_path)
            return journal
        else:
            conn.commit()

    journal.statut = StatutImport.TERMINE
    journal.nb_creations = rapport_preparation.nb_mises_a_jour
    journal.nb_ignorees = rapport_preparation.nb_ignorees
    journal.nb_rejetees = rapport_preparation.nb_rejetees
    journal.nb_erreurs = rapport_preparation.nb_erreurs
    journal.nb_avertissements = rapport_preparation.nb_avertissements
    journal.id = import_repository.enregistrer_import(journal, db_path=db_path)
    import_repository.enregistrer_anomalies(journal.id, rapport_preparation.toutes_anomalies, db_path=db_path)
    _journaliser_audit_import(journal, TypeActionAudit.IMPORT_DONNEES, db_path=db_path)
    return journal


# ---------------------------------------------------------------------
# Audit
# ---------------------------------------------------------------------

def _journaliser_audit_import(journal: ImportJournal, type_action: TypeActionAudit, db_path: DbPath = None) -> None:
    try:
        from database.repositories import audit_log_repository
        audit_log_repository.enregistrer(
            AuditLog(
                type_action=type_action, entite="import", entite_id=journal.id, utilisateur=journal.utilisateur,
                details=(
                    f"{journal.nom_fichier} ({journal.type_import.value}) — "
                    f"{journal.nb_lignes} ligne(s), {journal.nb_creations} création(s), "
                    f"{journal.nb_mises_a_jour} mise(s) à jour, {journal.nb_rejetees} rejet(s), "
                    f"{journal.nb_erreurs} erreur(s)"
                ),
            ),
            db_path=db_path,
        )
    except Exception as erreur:  # noqa: BLE001
        logger.warning("Impossible de journaliser l'import dans l'audit : %s", erreur)


# ---------------------------------------------------------------------
# Préparation — Import RÉMUNÉRATIONS (respecte le verrouillage des périodes)
# ---------------------------------------------------------------------

def preparer_import_remunerations(
    analyse: RapportAnalyse, df: pd.DataFrame, periode_id: int, db_path: DbPath = None
) -> RapportPreparation:
    """Même structure que preparer_import_heures, pour les 3 éléments de rémunération (section 8C)."""
    from utils.validators import valider_montant_fcfa

    periode = periode_repository.obtenir_par_id(periode_id, db_path=db_path)
    if periode is None:
        raise ImportServiceError(f"Aucune période avec l'id {periode_id}.")
    periode_modifiable = periode.statut == StatutPeriode.OUVERTE

    enseignants_existants = enseignant_repository.lister(inclure_inactifs=True, db_path=db_path)
    index_existants = {(_cle_normalisation(e.nom), _cle_normalisation(e.prenom)): e for e in enseignants_existants}

    champs_montants = ["prime_ap_pp", "surveillance_secretariat", "indemnite_suggestion_admin"]

    rapport = RapportPreparation()
    for position, row in enumerate(df.to_dict(orient="records"), start=1):
        ligne = LigneImport(numero_ligne=position, donnees=row)

        nom_brut = _valeur_normalisee(row, analyse.colonnes_reconnues, "nom")
        prenom_brut = _valeur_normalisee(row, analyse.colonnes_reconnues, "prenom")
        try:
            nom = valider_nom_ou_prenom(nom_brut, "nom")
            prenom = valider_nom_ou_prenom(prenom_brut, "prenom")
        except EnseignantValidationError as erreur:
            _ajouter_erreur(ligne, "nom", nom_brut, str(erreur))
            rapport.lignes.append(ligne)
            continue

        cle = (_cle_normalisation(nom), _cle_normalisation(prenom))
        enseignant = index_existants.get(cle)
        if enseignant is None:
            _ajouter_erreur(ligne, "nom", f"{nom} {prenom}", f"Enseignant inconnu : « {nom} {prenom} ».")
            rapport.lignes.append(ligne)
            continue
        if not enseignant.est_complet:
            _ajouter_erreur(ligne, "nom", f"{nom} {prenom}", message_fiche_incomplete(enseignant))
            rapport.lignes.append(ligne)
            continue

        if not periode_modifiable:
            _ajouter_erreur(
                ligne, "periode", periode.libelle,
                f"Cette période est {periode.statut.value.upper()}. Les données ne peuvent plus être modifiées.",
            )
            rapport.lignes.append(ligne)
            continue

        montants = {}
        erreur_montant = False
        for champ in champs_montants:
            valeur_brute = _valeur_normalisee(row, analyse.colonnes_reconnues, champ)
            if valeur_brute is None:
                montants[champ] = 0
                continue
            try:
                montants[champ] = valider_montant_fcfa(valeur_brute, champ)
            except ValueError as erreur:
                _ajouter_erreur(ligne, champ, valeur_brute, str(erreur))
                erreur_montant = True
        if erreur_montant:
            rapport.lignes.append(ligne)
            continue

        if all(v == 0 for v in montants.values()):
            _ajouter_avertissement(ligne, None, None, "Aucun montant renseigné sur cette ligne — ignorée.")
            ligne.action = ActionLigne.IGNOREE
            rapport.lignes.append(ligne)
            continue

        ligne.action = ActionLigne.MISE_A_JOUR
        ligne.cible_id = enseignant.id
        ligne.donnees = {"enseignant_id": enseignant.id, "montants": montants}
        rapport.lignes.append(ligne)

    return rapport


def executer_import_remunerations(
    rapport_preparation: RapportPreparation,
    nom_fichier: str,
    periode_id: int,
    utilisateur: Optional[str] = None,
    db_path: DbPath = None,
) -> ImportJournal:
    """Import transactionnel des rémunérations — mêmes garanties que les autres types."""
    from database.repositories import remuneration_repository
    from models.element_remuneration import ElementRemuneration
    from models.enums import TypeElementRemuneration

    journal = ImportJournal(
        nom_fichier=nom_fichier, type_import=TypeImport.REMUNERATIONS, statut=StatutImport.ECHEC,
        utilisateur=utilisateur, periode_id=periode_id, nb_lignes=rapport_preparation.nb_lignes,
    )
    correspondance_type = {
        "prime_ap_pp": TypeElementRemuneration.PRIME_AP_PP,
        "surveillance_secretariat": TypeElementRemuneration.SURVEILLANCE_SECRETARIAT,
        "indemnite_suggestion_admin": TypeElementRemuneration.INDEMNITE_SUGGESTION_ADMIN,
    }

    with get_connection(db_path) as conn:
        try:
            for ligne in rapport_preparation.lignes:
                if ligne.action != ActionLigne.MISE_A_JOUR:
                    continue
                for champ, montant in ligne.donnees["montants"].items():
                    element = ElementRemuneration(
                        enseignant_id=ligne.donnees["enseignant_id"], periode_id=periode_id,
                        type_element=correspondance_type[champ], montant=montant,
                    )
                    remuneration_repository.upsert(element, conn=conn)
        except Exception as erreur:  # noqa: BLE001
            conn.rollback()
            logger.error("Import rémunérations annulé (ROLLBACK) : %s", erreur)
            journal.nb_erreurs = rapport_preparation.nb_erreurs + 1
            journal.id = import_repository.enregistrer_import(journal, db_path=db_path)
            import_repository.enregistrer_anomalies(journal.id, rapport_preparation.toutes_anomalies, db_path=db_path)
            _journaliser_audit_import(journal, TypeActionAudit.IMPORT_DONNEES_ECHEC, db_path=db_path)
            return journal
        else:
            conn.commit()

    journal.statut = StatutImport.TERMINE
    journal.nb_creations = rapport_preparation.nb_mises_a_jour
    journal.nb_ignorees = rapport_preparation.nb_ignorees
    journal.nb_rejetees = rapport_preparation.nb_rejetees
    journal.nb_erreurs = rapport_preparation.nb_erreurs
    journal.nb_avertissements = rapport_preparation.nb_avertissements
    journal.id = import_repository.enregistrer_import(journal, db_path=db_path)
    import_repository.enregistrer_anomalies(journal.id, rapport_preparation.toutes_anomalies, db_path=db_path)
    _journaliser_audit_import(journal, TypeActionAudit.IMPORT_DONNEES, db_path=db_path)
    return journal


# ---------------------------------------------------------------------
# Préparation — Import RETENUES (respecte le verrouillage des périodes)
# ---------------------------------------------------------------------

def preparer_import_retenues(
    analyse: RapportAnalyse, df: pd.DataFrame, periode_id: int, db_path: DbPath = None
) -> RapportPreparation:
    """Même structure que preparer_import_remunerations, pour retenue_amicale/dette (section 8D)."""
    from utils.validators import valider_montant_fcfa

    periode = periode_repository.obtenir_par_id(periode_id, db_path=db_path)
    if periode is None:
        raise ImportServiceError(f"Aucune période avec l'id {periode_id}.")
    periode_modifiable = periode.statut == StatutPeriode.OUVERTE

    enseignants_existants = enseignant_repository.lister(inclure_inactifs=True, db_path=db_path)
    index_existants = {(_cle_normalisation(e.nom), _cle_normalisation(e.prenom)): e for e in enseignants_existants}

    champs_montants = ["retenue_amicale", "dette"]

    rapport = RapportPreparation()
    for position, row in enumerate(df.to_dict(orient="records"), start=1):
        ligne = LigneImport(numero_ligne=position, donnees=row)

        nom_brut = _valeur_normalisee(row, analyse.colonnes_reconnues, "nom")
        prenom_brut = _valeur_normalisee(row, analyse.colonnes_reconnues, "prenom")
        try:
            nom = valider_nom_ou_prenom(nom_brut, "nom")
            prenom = valider_nom_ou_prenom(prenom_brut, "prenom")
        except EnseignantValidationError as erreur:
            _ajouter_erreur(ligne, "nom", nom_brut, str(erreur))
            rapport.lignes.append(ligne)
            continue

        cle = (_cle_normalisation(nom), _cle_normalisation(prenom))
        enseignant = index_existants.get(cle)
        if enseignant is None:
            _ajouter_erreur(ligne, "nom", f"{nom} {prenom}", f"Enseignant inconnu : « {nom} {prenom} ».")
            rapport.lignes.append(ligne)
            continue
        if not enseignant.est_complet:
            _ajouter_erreur(ligne, "nom", f"{nom} {prenom}", message_fiche_incomplete(enseignant))
            rapport.lignes.append(ligne)
            continue

        if not periode_modifiable:
            _ajouter_erreur(
                ligne, "periode", periode.libelle,
                f"Cette période est {periode.statut.value.upper()}. Les données ne peuvent plus être modifiées.",
            )
            rapport.lignes.append(ligne)
            continue

        montants = {}
        erreur_montant = False
        for champ in champs_montants:
            valeur_brute = _valeur_normalisee(row, analyse.colonnes_reconnues, champ)
            if valeur_brute is None:
                montants[champ] = 0
                continue
            try:
                montants[champ] = valider_montant_fcfa(valeur_brute, champ)
            except ValueError as erreur:
                _ajouter_erreur(ligne, champ, valeur_brute, str(erreur))
                erreur_montant = True
        if erreur_montant:
            rapport.lignes.append(ligne)
            continue

        if all(v == 0 for v in montants.values()):
            _ajouter_avertissement(ligne, None, None, "Aucun montant renseigné sur cette ligne — ignorée.")
            ligne.action = ActionLigne.IGNOREE
            rapport.lignes.append(ligne)
            continue

        ligne.action = ActionLigne.MISE_A_JOUR
        ligne.cible_id = enseignant.id
        ligne.donnees = {"enseignant_id": enseignant.id, "montants": montants}
        rapport.lignes.append(ligne)

    return rapport


def executer_import_retenues(
    rapport_preparation: RapportPreparation,
    nom_fichier: str,
    periode_id: int,
    utilisateur: Optional[str] = None,
    db_path: DbPath = None,
) -> ImportJournal:
    """Import transactionnel des retenues — mêmes garanties que les autres types."""
    from database.repositories import retenue_repository
    from models.retenue import Retenue
    from models.enums import TypeRetenue

    journal = ImportJournal(
        nom_fichier=nom_fichier, type_import=TypeImport.RETENUES, statut=StatutImport.ECHEC,
        utilisateur=utilisateur, periode_id=periode_id, nb_lignes=rapport_preparation.nb_lignes,
    )
    correspondance_type = {"retenue_amicale": TypeRetenue.RETENUE_AMICALE, "dette": TypeRetenue.DETTE}

    with get_connection(db_path) as conn:
        try:
            for ligne in rapport_preparation.lignes:
                if ligne.action != ActionLigne.MISE_A_JOUR:
                    continue
                for champ, montant in ligne.donnees["montants"].items():
                    retenue = Retenue(
                        enseignant_id=ligne.donnees["enseignant_id"], periode_id=periode_id,
                        type_retenue=correspondance_type[champ], montant=montant,
                    )
                    retenue_repository.upsert(retenue, conn=conn)
        except Exception as erreur:  # noqa: BLE001
            conn.rollback()
            logger.error("Import retenues annulé (ROLLBACK) : %s", erreur)
            journal.nb_erreurs = rapport_preparation.nb_erreurs + 1
            journal.id = import_repository.enregistrer_import(journal, db_path=db_path)
            import_repository.enregistrer_anomalies(journal.id, rapport_preparation.toutes_anomalies, db_path=db_path)
            _journaliser_audit_import(journal, TypeActionAudit.IMPORT_DONNEES_ECHEC, db_path=db_path)
            return journal
        else:
            conn.commit()

    journal.statut = StatutImport.TERMINE
    journal.nb_creations = rapport_preparation.nb_mises_a_jour
    journal.nb_ignorees = rapport_preparation.nb_ignorees
    journal.nb_rejetees = rapport_preparation.nb_rejetees
    journal.nb_erreurs = rapport_preparation.nb_erreurs
    journal.nb_avertissements = rapport_preparation.nb_avertissements
    journal.id = import_repository.enregistrer_import(journal, db_path=db_path)
    import_repository.enregistrer_anomalies(journal.id, rapport_preparation.toutes_anomalies, db_path=db_path)
    _journaliser_audit_import(journal, TypeActionAudit.IMPORT_DONNEES, db_path=db_path)
    return journal
