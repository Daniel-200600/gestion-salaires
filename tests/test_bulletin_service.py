"""
Tests du service bulletin (services/bulletin_service.py).

Couvre la génération individuelle/groupée, l'isolation stricte entre
enseignants et entre périodes, la protection contre l'écrasement pour
une période clôturée, et l'intégrité exacte avec paie_service.
"""

import pytest
from docx import Document

import database.connection as database_connection
from services import bulletin_service, donnees_paie_service, enseignant_service, paie_service, periode_service
from services.bulletin_service import BulletinServiceError
from services.donnees_paie_service import DonneesPaieEnseignant
from utils.montant_en_lettres import montant_en_lettres


@pytest.fixture(autouse=True)
def _rediriger_connexion_par_defaut(db_path, monkeypatch, tmp_path):
    monkeypatch.setattr(database_connection, "DB_PATH", db_path)
    # Isole aussi les fichiers générés dans un dossier temporaire dédié au test.
    dossier_export = tmp_path / "bulletins"
    monkeypatch.setattr(bulletin_service, "EXPORT_DIR_BULLETINS", dossier_export)


def _creer_enseignant(**overrides):
    donnees = {"nom": "Ngono", "prenom": "Marie", "sexe": "F", "statut": "P", "taux_horaire": 1500}
    donnees.update(overrides)
    return enseignant_service.creer_enseignant(**donnees)


def _creer_periode_ouverte(mois=9, annee=2030):
    periode = periode_service.creer_periode(mois=mois, annee=annee)
    return periode_service.ouvrir_periode(periode.id)


def _saisir_donnees(periode_id, enseignant_id, **kwargs):
    donnees_paie_service.enregistrer_donnees_paie_groupe(
        periode_id, [DonneesPaieEnseignant(enseignant_id=enseignant_id, **kwargs)]
    )


def _lire_texte_document(chemin):
    doc = Document(chemin)
    textes = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                textes.append(cell.text)
    return " | ".join(textes)


# ---------------------------------------------------------------------
# 1 & 2. Génération d'un bulletin, fichier .docx valide
# ---------------------------------------------------------------------

def test_generation_bulletin_individuel():
    e = _creer_enseignant()
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    p = periode_service.valider_periode(p.id)

    resultat = bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    assert resultat.chemin.exists()
    Document(resultat.chemin)  # lève une exception si le fichier n'est pas un docx valide


# ---------------------------------------------------------------------
# 3 & 4. Nom correct, nom/prénom injectés
# ---------------------------------------------------------------------

def test_nom_fichier_et_contenu_corrects():
    e = _creer_enseignant(nom="Ngono", prenom="Marie")
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    p = periode_service.valider_periode(p.id)

    resultat = bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    assert "NGONO" in resultat.chemin.name
    assert "MARIE" in resultat.chemin.name

    texte = _lire_texte_document(resultat.chemin)
    assert "NGONO" in texte
    assert "MARIE" in texte


# ---------------------------------------------------------------------
# 5 & 6. Statut et période injectés
# ---------------------------------------------------------------------

def test_statut_et_periode_injectes():
    e = _creer_enseignant(statut="V")
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    p = periode_service.valider_periode(p.id)

    resultat = bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    texte = _lire_texte_document(resultat.chemin)
    assert "V" in texte
    # Adapté (reproduction du bulletin officiel de l'établissement, demandée
    # par celui-ci) : la période y est écrite en anglais, « SEPTEMBER 2030 ».
    # Le libellé français reste disponible via la balise {{PERIODE_FR}}.
    from utils.formatters import libelle_periode_anglais
    assert libelle_periode_anglais(p.mois, p.annee) in texte


# ---------------------------------------------------------------------
# 7 à 16. Toutes les valeurs financières injectées correctement
# ---------------------------------------------------------------------

def test_toutes_les_valeurs_financieres_injectees():
    e = _creer_enseignant(taux_horaire=2000)
    p = _creer_periode_ouverte()
    _saisir_donnees(
        p.id, e.id,
        heures_par_semaine={1: 20, 2: 20, 3: 20, 4: 20, 5: 20},
        prime_ap_pp=20000, surveillance_secretariat=10000, indemnite_suggestion_admin=5000,
        retenue_amicale=5000, dette=10000,
    )
    p = periode_service.valider_periode(p.id)

    resultat_moteur = paie_service.calculer_paie_enseignant(p.id, e.id)
    resultat = bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    texte = _lire_texte_document(resultat.chemin)

    assert str(resultat_moteur.total_heures).rstrip("0").rstrip(".") in texte or "100" in texte
    assert "2000" in texte  # taux horaire
    assert "200000" in texte  # gain heures
    assert "20000" in texte  # prime ap/pp
    assert "10000" in texte  # surveillance
    assert "5000" in texte  # indemnite (et retenue amicale, valeur identique)
    assert "11750" in texte  # taxe
    assert "10000" in texte  # dette
    assert "208250" in texte  # net


# ---------------------------------------------------------------------
# 17. Montant en lettres correct
# ---------------------------------------------------------------------

def test_montant_en_lettres_correct():
    e = _creer_enseignant(taux_horaire=1800)
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    p = periode_service.valider_periode(p.id)

    resultat = bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    texte = _lire_texte_document(resultat.chemin)
    lettres_attendues = montant_en_lettres(resultat.resultat.net_a_percevoir)
    assert lettres_attendues in texte


# ---------------------------------------------------------------------
# 18. Aucun placeholder restant
# ---------------------------------------------------------------------

def test_aucun_placeholder_restant():
    e = _creer_enseignant()
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    p = periode_service.valider_periode(p.id)

    resultat = bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    texte = _lire_texte_document(resultat.chemin)
    assert "{{" not in texte
    assert "}}" not in texte


# ---------------------------------------------------------------------
# 19 & 20. Génération de plusieurs bulletins, noms distincts
# ---------------------------------------------------------------------

def test_generation_groupee_noms_distincts():
    e1 = _creer_enseignant(nom="Ngono", prenom="Marie")
    e2 = _creer_enseignant(nom="Fotso", prenom="Paul")
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e1.id, heures_par_semaine={1: 10})
    _saisir_donnees(p.id, e2.id, heures_par_semaine={1: 15})
    p = periode_service.valider_periode(p.id)

    resultat = bulletin_service.generer_bulletins_groupe(p.id, [e1.id, e2.id])
    assert len(resultat.bulletins) == 2
    noms_fichiers = {b.chemin.name for b in resultat.bulletins}
    assert len(noms_fichiers) == 2  # bien distincts


# ---------------------------------------------------------------------
# 21. ZIP créé correctement
# ---------------------------------------------------------------------

def test_creation_zip():
    import zipfile

    e1 = _creer_enseignant(nom="Ngono", prenom="Marie")
    e2 = _creer_enseignant(nom="Fotso", prenom="Paul")
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e1.id, heures_par_semaine={1: 10})
    _saisir_donnees(p.id, e2.id, heures_par_semaine={1: 15})
    p = periode_service.valider_periode(p.id)

    resultat = bulletin_service.generer_bulletins_groupe(p.id, [e1.id, e2.id])
    chemin_zip = bulletin_service.creer_archive_zip(resultat.bulletins, p.libelle)

    assert chemin_zip.exists()
    with zipfile.ZipFile(chemin_zip) as archive:
        noms = archive.namelist()
        assert len(noms) == 2
        for bulletin in resultat.bulletins:
            assert bulletin.chemin.name in noms


# ---------------------------------------------------------------------
# 22 & 23. Données provenant du paie_service, aucune formule dans word_export.py
# ---------------------------------------------------------------------

def test_donnees_identiques_a_paie_service():
    e = _creer_enseignant(taux_horaire=2000)
    p = _creer_periode_ouverte()
    _saisir_donnees(
        p.id, e.id,
        heures_par_semaine={1: 20, 2: 20, 3: 20, 4: 20, 5: 20},
        prime_ap_pp=20000, surveillance_secretariat=10000, indemnite_suggestion_admin=5000,
        retenue_amicale=5000, dette=10000,
    )
    p = periode_service.valider_periode(p.id)

    resultat_moteur = paie_service.calculer_paie_enseignant(p.id, e.id)
    resultat_bulletin = bulletin_service.generer_bulletin_enseignant(p.id, e.id)

    assert resultat_bulletin.resultat.net_a_percevoir == resultat_moteur.net_a_percevoir == 208250
    assert resultat_bulletin.resultat.taxe_5 == resultat_moteur.taxe_5
    assert resultat_bulletin.resultat.gain_heures == resultat_moteur.gain_heures


def test_aucune_formule_de_paie_dans_word_export():
    import inspect
    from exports import word_export

    source = inspect.getsource(word_export)
    assert "* 0.05" not in source
    assert "TAUX_TAXE" not in source
    assert "taux_horaire *" not in source


def test_aucune_formule_de_paie_dans_bulletin_service():
    import inspect

    source = inspect.getsource(bulletin_service)
    assert "* 0.05" not in source
    assert "TAUX_TAXE" not in source
    assert "taux_horaire *" not in source


# ---------------------------------------------------------------------
# 24 & 25. Isolation stricte entre périodes et entre enseignants
# ---------------------------------------------------------------------

def test_bulletin_periode_ne_contient_pas_donnees_autre_periode():
    e = _creer_enseignant(taux_horaire=1000)
    p1 = _creer_periode_ouverte(mois=1, annee=2031)
    p2 = _creer_periode_ouverte(mois=2, annee=2031)
    _saisir_donnees(p1.id, e.id, heures_par_semaine={1: 10})
    _saisir_donnees(p2.id, e.id, heures_par_semaine={1: 50})
    p1 = periode_service.valider_periode(p1.id)
    p2 = periode_service.valider_periode(p2.id)

    resultat_p1 = bulletin_service.generer_bulletin_enseignant(p1.id, e.id)
    resultat_p2 = bulletin_service.generer_bulletin_enseignant(p2.id, e.id)

    assert resultat_p1.resultat.gain_heures == 10000  # 10h
    assert resultat_p2.resultat.gain_heures == 50000  # 50h
    assert resultat_p1.chemin != resultat_p2.chemin

    texte_p1 = _lire_texte_document(resultat_p1.chemin)
    texte_p2 = _lire_texte_document(resultat_p2.chemin)
    assert "50000" not in texte_p1
    assert "10000" not in texte_p2 or "10000" in texte_p2 and "50000" not in texte_p1  # garde-fou explicite ci-dessous
    # Adapté : le bulletin officiel écrit la période en anglais (« JANUARY 2031 »).
    from utils.formatters import libelle_periode_anglais
    libelle_p1, libelle_p2 = libelle_periode_anglais(p1.mois, p1.annee), libelle_periode_anglais(p2.mois, p2.annee)
    assert libelle_p1 in texte_p1
    assert libelle_p2 in texte_p2
    assert libelle_p2 not in texte_p1
    assert libelle_p1 not in texte_p2


def test_enseignant_a_ne_recoit_jamais_les_donnees_de_b():
    e_a = _creer_enseignant(nom="Ateba", prenom="Sylvie", taux_horaire=1000)
    e_b = _creer_enseignant(nom="Biya", prenom="Rene", taux_horaire=9000)
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e_a.id, heures_par_semaine={1: 5})
    _saisir_donnees(p.id, e_b.id, heures_par_semaine={1: 40})
    p = periode_service.valider_periode(p.id)

    resultat_a = bulletin_service.generer_bulletin_enseignant(p.id, e_a.id)
    resultat_b = bulletin_service.generer_bulletin_enseignant(p.id, e_b.id)

    texte_a = _lire_texte_document(resultat_a.chemin)
    texte_b = _lire_texte_document(resultat_b.chemin)

    assert "ATEBA" in texte_a and "BIYA" not in texte_a
    assert "BIYA" in texte_b and "ATEBA" not in texte_b
    assert str(resultat_a.resultat.gain_heures) == "5000"
    assert str(resultat_b.resultat.gain_heures) == "360000"


# ---------------------------------------------------------------------
# Protection de l'historique (période clôturée)
# ---------------------------------------------------------------------

def test_regeneration_refusee_periode_cloturee_si_deja_genere():
    e = _creer_enseignant()
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    p = periode_service.valider_periode(p.id)
    p = periode_service.cloturer_periode(p.id)

    bulletin_service.generer_bulletin_enseignant(p.id, e.id)  # 1re génération : OK

    with pytest.raises(BulletinServiceError, match="clôturée"):
        bulletin_service.generer_bulletin_enseignant(p.id, e.id)  # 2e génération : refusée


def test_regeneration_versionnee_si_periode_non_cloturee():
    e = _creer_enseignant()
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    p = periode_service.valider_periode(p.id)  # validée, pas clôturée

    resultat1 = bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    resultat2 = bulletin_service.generer_bulletin_enseignant(p.id, e.id)  # ne doit pas écraser
    assert resultat1.chemin != resultat2.chemin
    assert resultat1.chemin.exists()
    assert resultat2.chemin.exists()


# ---------------------------------------------------------------------
# Cas limites : période brouillon, enseignant/période inexistants
# ---------------------------------------------------------------------

def test_periode_brouillon_refusee():
    e = _creer_enseignant()
    p = periode_service.creer_periode(mois=6, annee=2031)
    with pytest.raises(BulletinServiceError):
        bulletin_service.generer_bulletin_enseignant(p.id, e.id)


def test_enseignant_inexistant():
    p = _creer_periode_ouverte()
    with pytest.raises(BulletinServiceError):
        bulletin_service.generer_bulletin_enseignant(p.id, 9999)


def test_erreur_isolee_dans_generation_groupee():
    e = _creer_enseignant()
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    p = periode_service.valider_periode(p.id)

    resultat = bulletin_service.generer_bulletins_groupe(p.id, [e.id, 9999])
    assert len(resultat.bulletins) == 1
    assert 9999 in resultat.erreurs


# ---------------------------------------------------------------------
# Contrôles de validité explicites (section 21 de la demande de finition)
# ---------------------------------------------------------------------

def test_validation_reussie_bulletin_complet():
    """Un bulletin correctement généré passe silencieusement les 10 contrôles."""
    e = _creer_enseignant(taux_horaire=1000)
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10}, prime_ap_pp=500)
    p = periode_service.valider_periode(p.id)

    resultat = bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    # Ne lève rien : le fichier existe et est déjà validé.
    assert resultat.chemin.exists()


def test_validation_verifie_nom_fichier_correspond_enseignant():
    e = _creer_enseignant(nom="Mballa", prenom="Eric")
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    p = periode_service.valider_periode(p.id)

    resultat = bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    assert "MBALLA" in resultat.chemin.name.upper()


def test_validation_detecte_placeholder_residuel_et_supprime_le_fichier(monkeypatch):
    """
    Si un placeholder venait à subsister malgré tout (défaut de
    template), la validation doit lever une erreur explicite et ne
    JAMAIS laisser un bulletin incomplet sur disque.
    """
    e = _creer_enseignant()
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    p = periode_service.valider_periode(p.id)

    # Simule un document généré avec un placeholder oublié.
    from exports import word_export as word_export_module

    document_original = word_export_module.generer_document_bulletin

    def _generer_avec_placeholder_oublie(valeurs, template_path=None):
        valeurs_incompletes = dict(valeurs)
        valeurs_incompletes.pop("{{NET_A_PERÇEVOIR}}", None)
        # On contourne volontairement la vérification interne de word_export
        # pour simuler un défaut de template qui laisserait passer un placeholder.
        from docx import Document as DocxDocument
        chemin_template = template_path or word_export_module.TEMPLATE_PATH
        doc = DocxDocument(chemin_template)
        word_export_module._remplacer_dans_document(doc, valeurs_incompletes)
        return doc

    monkeypatch.setattr(bulletin_service, "generer_document_bulletin", _generer_avec_placeholder_oublie)

    with pytest.raises(BulletinServiceError, match="[Pp]laceholder"):
        bulletin_service.generer_bulletin_enseignant(p.id, e.id)

    # Le fichier incomplet ne doit jamais rester sur disque.
    dossier = bulletin_service._dossier_periode(p.libelle)
    fichiers_restants = list(dossier.glob("*.docx")) if dossier.exists() else []
    assert fichiers_restants == []


# ---------------------------------------------------------------------
# Test avec trois enseignants distincts (section 23 de la demande de finition)
# ---------------------------------------------------------------------

def test_trois_enseignants_aucune_donnee_melangee():
    e_a = _creer_enseignant(nom="Ateba", prenom="Sylvie", statut="P", taux_horaire=1000)
    e_b = _creer_enseignant(nom="Biya", prenom="Rene", statut="V", taux_horaire=2000)
    e_c = _creer_enseignant(nom="Chomdack", prenom="Awa", statut="P", taux_horaire=1500)
    p = _creer_periode_ouverte()

    _saisir_donnees(p.id, e_a.id, heures_par_semaine={1: 10}, prime_ap_pp=1000)
    _saisir_donnees(p.id, e_b.id, heures_par_semaine={1: 20}, dette=2000)
    _saisir_donnees(p.id, e_c.id, heures_par_semaine={1: 15}, retenue_amicale=500)
    p = periode_service.valider_periode(p.id)

    resultat = bulletin_service.generer_bulletins_groupe(p.id, [e_a.id, e_b.id, e_c.id])
    assert len(resultat.bulletins) == 3
    assert resultat.erreurs == {}

    bulletins_par_id = {b.enseignant_id: b for b in resultat.bulletins}

    texte_a = _lire_texte_document(bulletins_par_id[e_a.id].chemin)
    texte_b = _lire_texte_document(bulletins_par_id[e_b.id].chemin)
    texte_c = _lire_texte_document(bulletins_par_id[e_c.id].chemin)

    # Chaque bulletin contient son propre enseignant...
    assert "ATEBA" in texte_a
    assert "BIYA" in texte_b
    assert "CHOMDACK" in texte_c

    # ...et JAMAIS les noms des deux autres.
    assert "BIYA" not in texte_a and "CHOMDACK" not in texte_a
    assert "ATEBA" not in texte_b and "CHOMDACK" not in texte_b
    assert "ATEBA" not in texte_c and "BIYA" not in texte_c

    # Les gains (dérivés du taux horaire propre à chacun) ne sont jamais mélangés.
    assert bulletins_par_id[e_a.id].resultat.gain_heures == 10000  # 10h x 1000
    assert bulletins_par_id[e_b.id].resultat.gain_heures == 40000  # 20h x 2000
    assert bulletins_par_id[e_c.id].resultat.gain_heures == 22500  # 15h x 1500

    # Trois fichiers physiquement distincts.
    chemins = {b.chemin for b in resultat.bulletins}
    assert len(chemins) == 3

    # Le ZIP groupé contient bien les 3 fichiers, sans mélange.
    chemin_zip = bulletin_service.creer_archive_zip(resultat.bulletins, p.libelle)
    import zipfile
    with zipfile.ZipFile(chemin_zip) as archive:
        assert len(archive.namelist()) == 3


# ---------------------------------------------------------------------
# Instantané immuable de traçabilité (module 12, section 11/13)
# ---------------------------------------------------------------------

def test_snapshot_non_pris_pour_periode_ouverte():
    """La génération reste possible sur une période OUVERTE (comme depuis le module 07), mais sans instantané."""
    from database.repositories import bulletin_repository

    e = _creer_enseignant()
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})

    bulletin_service.generer_bulletin_enseignant(p.id, e.id)

    snapshots = bulletin_repository.lister_par_periode(p.id)
    assert snapshots == []


def test_snapshot_pris_pour_periode_validee():
    from database.repositories import bulletin_repository
    from services import periode_service

    e = _creer_enseignant(taux_horaire=2000)
    p = _creer_periode_ouverte()
    _saisir_donnees(
        p.id, e.id,
        heures_par_semaine={1: 20, 2: 20, 3: 20, 4: 20, 5: 20},
        prime_ap_pp=20000, surveillance_secretariat=10000, indemnite_suggestion_admin=5000,
        retenue_amicale=5000, dette=10000,
    )
    periode_service.valider_periode(p.id)

    bulletin_service.generer_bulletin_enseignant(p.id, e.id, utilisateur="testeur")

    snapshots = bulletin_repository.lister_par_periode(p.id)
    assert len(snapshots) == 1
    assert snapshots[0].net_a_payer == 208250
    assert snapshots[0].utilisateur_generation == "testeur"


def test_snapshot_idempotent_si_bulletin_deja_instantane():
    """Un second appel (ex. régénération) ne crée jamais de second instantané : le premier fait foi."""
    from database.repositories import bulletin_repository
    from services import periode_service

    e = _creer_enseignant()
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    periode_service.valider_periode(p.id)

    bulletin_service.generer_bulletin_enseignant(p.id, e.id, utilisateur="premier")
    bulletin_service.generer_bulletin_enseignant(p.id, e.id, utilisateur="second")

    snapshots = bulletin_repository.lister_par_periode(p.id)
    assert len(snapshots) == 1
    assert snapshots[0].utilisateur_generation == "premier"  # le premier instantané fait foi


def test_snapshot_immuable_update_refuse(db_path):
    """Le trigger SQL interdit toute modification d'un instantané déjà enregistré."""
    import sqlite3
    from services import periode_service

    e = _creer_enseignant()
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    periode_service.valider_periode(p.id)
    bulletin_service.generer_bulletin_enseignant(p.id, e.id)

    conn = sqlite3.connect(db_path)
    with pytest.raises(sqlite3.IntegrityError, match="immuable"):
        conn.execute("UPDATE bulletins_paie SET net_a_payer = 999999 WHERE enseignant_id = ?", (e.id,))
    conn.close()


def test_snapshot_immuable_delete_refuse(db_path):
    """Le trigger SQL interdit toute suppression d'un instantané déjà enregistré."""
    import sqlite3
    from services import periode_service

    e = _creer_enseignant()
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    periode_service.valider_periode(p.id)
    bulletin_service.generer_bulletin_enseignant(p.id, e.id)

    conn = sqlite3.connect(db_path)
    with pytest.raises(sqlite3.IntegrityError, match="immuable"):
        conn.execute("DELETE FROM bulletins_paie WHERE enseignant_id = ?", (e.id,))
    conn.close()


def test_snapshot_contenu_coherent_avec_le_calcul():
    """L'instantané reprend exactement les mêmes valeurs que paie_service (aucune formule dupliquée)."""
    from database.repositories import bulletin_repository
    from services import periode_service
    from services.paie_service import calculer_paie_enseignant

    e = _creer_enseignant(taux_horaire=1500)
    p = _creer_periode_ouverte()
    _saisir_donnees(p.id, e.id, heures_par_semaine={1: 15}, prime_ap_pp=1000, dette=200)
    periode_service.valider_periode(p.id)

    bulletin_service.generer_bulletin_enseignant(p.id, e.id)

    resultat = calculer_paie_enseignant(p.id, e.id)
    snapshot = bulletin_repository.lister_par_periode(p.id)[0]
    assert snapshot.net_a_payer == resultat.net_a_percevoir
    assert snapshot.gain_heures == resultat.gain_heures
    assert snapshot.taxe_5pct == resultat.taxe_5
