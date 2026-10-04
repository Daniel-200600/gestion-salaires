"""
Tests du générateur Excel (exports/excel_export.py).

Vérifie la structure du classeur (3 feuilles, colonnes dans l'ordre),
le contenu (noms/prénoms/statuts/totaux/semaines), le nom de fichier
dynamique et le nettoyage des caractères interdits Windows, et surtout
l'INTÉGRITÉ EXACTE entre les résultats du moteur de paie et les
valeurs écrites dans le fichier Excel.
"""

import shutil
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

import database.connection as database_connection
from exports import excel_export
from services import comptabilite_service, donnees_paie_service, enseignant_service, periode_service
from services.donnees_paie_service import DonneesPaieEnseignant

# Cas de référence officiel : vacataire, taxe de 5 %.
pytestmark = pytest.mark.taxe_historique


@pytest.fixture(autouse=True)
def _rediriger_connexion_par_defaut(db_path, monkeypatch):
    monkeypatch.setattr(database_connection, "DB_PATH", db_path)


@pytest.fixture
def dossier_export(tmp_path):
    dossier = tmp_path / "exports_test"
    yield dossier
    shutil.rmtree(dossier, ignore_errors=True)


def _creer_enseignant(**overrides):
    donnees = {"nom": "Sissoko", "prenom": "Modibo", "sexe": "M", "statut": "V", "taux_horaire": 1000}
    donnees.update(overrides)
    return enseignant_service.creer_enseignant(**donnees)


def _creer_periode_ouverte(mois=3, annee=2030):
    periode = periode_service.creer_periode(mois=mois, annee=annee)
    return periode_service.ouvrir_periode(periode.id)


def _saisir_donnees(periode_id, enseignant_id, **kwargs):
    donnees_paie_service.enregistrer_donnees_paie_groupe(
        periode_id, [DonneesPaieEnseignant(enseignant_id=enseignant_id, **kwargs)]
    )


def _trouver_ligne_entete(feuille, premiere_colonne_attendue: str) -> int:
    """Localise la ligne d'en-tête de colonnes (celle dont la 1re cellule correspond)."""
    for ligne in range(1, 20):
        if feuille.cell(row=ligne, column=1).value == premiere_colonne_attendue:
            return ligne
    raise AssertionError(f"En-tête de colonnes '{premiere_colonne_attendue}' introuvable.")


# ---------------------------------------------------------------------
# Fixture d'un état comptable simple, réutilisée par plusieurs tests
# ---------------------------------------------------------------------

def _etat_exemple_enonce():
    """Reproduit exactement l'exemple chiffré de l'énoncé (net attendu = 208250)."""
    e = _creer_enseignant(nom="Kamgang", prenom="Paul", statut="V", taux_horaire=2000)
    p = _creer_periode_ouverte()
    _saisir_donnees(
        p.id, e.id,
        heures_par_semaine={1: 20, 2: 20, 3: 20, 4: 20, 5: 20},
        prime_ap_pp=20000, surveillance_secretariat=10000, indemnite_suggestion_admin=5000,
        retenue_amicale=5000, dette=10000,
    )
    p = periode_service.valider_periode(p.id)
    return comptabilite_service.preparer_etat_comptable(p.id)


# ---------------------------------------------------------------------
# 9 & 10. Fichier créé correctement, .xlsx valide
# ---------------------------------------------------------------------

def test_fichier_cree_et_valide(dossier_export):
    etat = _etat_exemple_enonce()
    chemin = excel_export.generer_fichier_excel(etat, dossier=dossier_export)

    assert chemin.exists()
    assert chemin.suffix == ".xlsx"
    classeur = load_workbook(chemin)  # lève une exception si le fichier n'est pas un xlsx valide
    assert classeur is not None


# ---------------------------------------------------------------------
# 11, 12, 13. Les 3 feuilles attendues sont présentes
# ---------------------------------------------------------------------

def test_trois_feuilles_presentes(dossier_export):
    etat = _etat_exemple_enonce()
    chemin = excel_export.generer_fichier_excel(etat, dossier=dossier_export)
    classeur = load_workbook(chemin)
    assert classeur.sheetnames == ["État comptable", "Synthèse", "Détail heures"]


# ---------------------------------------------------------------------
# 14. Colonnes présentes dans le bon ordre
# ---------------------------------------------------------------------

def test_colonnes_etat_comptable_ordre_correct(dossier_export):
    etat = _etat_exemple_enonce()
    chemin = excel_export.generer_fichier_excel(etat, dossier=dossier_export)
    classeur = load_workbook(chemin)
    feuille = classeur["État comptable"]

    ligne_entete = _trouver_ligne_entete(feuille, "N°")
    colonnes = [feuille.cell(row=ligne_entete, column=c).value for c in range(1, 17)]
    assert colonnes == excel_export.COLONNES_ETAT_COMPTABLE


def test_colonnes_detail_heures_ordre_correct(dossier_export):
    etat = _etat_exemple_enonce()
    chemin = excel_export.generer_fichier_excel(etat, dossier=dossier_export)
    classeur = load_workbook(chemin)
    feuille = classeur["Détail heures"]

    ligne_entete = _trouver_ligne_entete(feuille, "N°")
    colonnes = [feuille.cell(row=ligne_entete, column=c).value for c in range(1, 11)]
    assert colonnes == excel_export.COLONNES_DETAIL_HEURES


# ---------------------------------------------------------------------
# 15. Nombre de lignes correct
# ---------------------------------------------------------------------

def test_nombre_de_lignes_correct(dossier_export):
    for i in range(3):
        _creer_enseignant(nom=f"Ens{i}", prenom=f"P{i}")
    # Réutilise une période + saisit pour les 3 nouveaux enseignants
    from database.repositories import enseignant_repository
    enseignants = [e for e in enseignant_repository.lister(inclure_inactifs=True)]
    p = _creer_periode_ouverte(mois=4)
    for e in enseignants:
        _saisir_donnees(p.id, e.id, heures_par_semaine={1: 10})
    p = periode_service.valider_periode(p.id)

    etat = comptabilite_service.preparer_etat_comptable(p.id)
    chemin = excel_export.generer_fichier_excel(etat, dossier=dossier_export)
    classeur = load_workbook(chemin)
    feuille = classeur["État comptable"]

    ligne_entete = _trouver_ligne_entete(feuille, "N°")
    # +1 ligne TOTAL GÉNÉRAL après les N lignes de données
    nb_lignes_donnees = len(etat.resultats)
    ligne_total = ligne_entete + nb_lignes_donnees + 1
    assert feuille.cell(row=ligne_total, column=2).value == "TOTAL GÉNÉRAL"


# ---------------------------------------------------------------------
# 16 & 17. Noms/prénoms et statuts corrects
# ---------------------------------------------------------------------

def test_noms_prenoms_statuts_corrects(dossier_export):
    etat = _etat_exemple_enonce()
    chemin = excel_export.generer_fichier_excel(etat, dossier=dossier_export)
    classeur = load_workbook(chemin)
    feuille = classeur["État comptable"]

    ligne_entete = _trouver_ligne_entete(feuille, "N°")
    ligne_donnees = ligne_entete + 1

    assert feuille.cell(row=ligne_donnees, column=2).value == "Kamgang"
    assert feuille.cell(row=ligne_donnees, column=3).value == "Paul"
    assert feuille.cell(row=ligne_donnees, column=4).value == "Masculin"
    assert feuille.cell(row=ligne_donnees, column=5).value == "Vacataire"


# ---------------------------------------------------------------------
# 18 & 19. Totaux corrects, total net correct
# ---------------------------------------------------------------------

def test_totaux_corrects(dossier_export):
    etat = _etat_exemple_enonce()
    chemin = excel_export.generer_fichier_excel(etat, dossier=dossier_export)
    classeur = load_workbook(chemin)
    feuille = classeur["État comptable"]

    ligne_entete = _trouver_ligne_entete(feuille, "N°")
    ligne_total = ligne_entete + len(etat.resultats) + 1

    assert feuille.cell(row=ligne_total, column=6).value == 100  # total heures
    assert feuille.cell(row=ligne_total, column=8).value == 200000  # total gain heures
    assert feuille.cell(row=ligne_total, column=13).value == 11750  # total taxe
    assert feuille.cell(row=ligne_total, column=16).value == 208250  # total net à percevoir


def test_total_net_correct_synthese(dossier_export):
    etat = _etat_exemple_enonce()
    chemin = excel_export.generer_fichier_excel(etat, dossier=dossier_export)
    classeur = load_workbook(chemin)
    feuille = classeur["Synthèse"]

    valeurs = [feuille.cell(row=r, column=1).value for r in range(1, 20)]
    ligne_net = next(r for r in range(1, 20) if feuille.cell(row=r, column=1).value == "Total net à percevoir")
    assert feuille.cell(row=ligne_net, column=2).value == 208250


# ---------------------------------------------------------------------
# 20. Données des semaines correctes (feuille Détail heures)
# ---------------------------------------------------------------------

def test_donnees_semaines_correctes(dossier_export):
    etat = _etat_exemple_enonce()
    chemin = excel_export.generer_fichier_excel(etat, dossier=dossier_export)
    classeur = load_workbook(chemin)
    feuille = classeur["Détail heures"]

    ligne_entete = _trouver_ligne_entete(feuille, "N°")
    ligne_donnees = ligne_entete + 1
    semaines = [feuille.cell(row=ligne_donnees, column=c).value for c in range(5, 10)]
    assert semaines == [20, 20, 20, 20, 20]
    assert feuille.cell(row=ligne_donnees, column=10).value == 100  # total heures


# ---------------------------------------------------------------------
# 21 & 22. Nom du fichier correct, caractères Windows interdits nettoyés
# ---------------------------------------------------------------------

def test_nom_fichier_dynamique_correct():
    etat = _etat_exemple_enonce()
    nom = excel_export.generer_nom_fichier(etat)
    assert nom == "Etat_comptable_Mars_2030.xlsx"


def test_nom_fichier_caracteres_interdits_nettoyes():
    etat = _etat_exemple_enonce()
    etat.periode.libelle = 'Test: <Période> "spéciale"/2030'
    nom = excel_export.generer_nom_fichier(etat)
    for caractere_interdit in '<>:"/\\|?*':
        assert caractere_interdit not in nom


def test_ne_jamais_ecraser_fichier_existant(dossier_export):
    etat = _etat_exemple_enonce()
    chemin1 = excel_export.generer_fichier_excel(etat, dossier=dossier_export)
    chemin2 = excel_export.generer_fichier_excel(etat, dossier=dossier_export)
    assert chemin1 != chemin2
    assert chemin1.exists()
    assert chemin2.exists()


# ---------------------------------------------------------------------
# Test d'intégrité — le plus important : Net application == Net Excel
# ---------------------------------------------------------------------

def test_integrite_exacte_net_application_egale_net_excel(dossier_export):
    """
    Cas de l'énoncé : 100h × 2000 FCFA, primes 20000/10000/5000,
    retenues 5000/10000 -> net attendu = 208250 FCFA. Le fichier Excel
    doit contenir EXACTEMENT cette même valeur, identique à celle du
    moteur de paie (aucune deuxième implémentation de la formule).
    """
    etat = _etat_exemple_enonce()
    resultat_moteur = etat.resultats[0]
    assert resultat_moteur.net_a_percevoir == 208250  # calcul manuel attendu par l'énoncé

    chemin = excel_export.generer_fichier_excel(etat, dossier=dossier_export)
    classeur = load_workbook(chemin)
    feuille = classeur["État comptable"]
    ligne_entete = _trouver_ligne_entete(feuille, "N°")
    ligne_donnees = ligne_entete + 1

    valeurs_excel = {
        "total_heures": feuille.cell(row=ligne_donnees, column=6).value,
        "taux_horaire": feuille.cell(row=ligne_donnees, column=7).value,
        "gain_heures": feuille.cell(row=ligne_donnees, column=8).value,
        "prime_ap_pp": feuille.cell(row=ligne_donnees, column=9).value,
        "surveillance_secretariat": feuille.cell(row=ligne_donnees, column=10).value,
        "indemnite_suggestion_admin": feuille.cell(row=ligne_donnees, column=11).value,
        "base_taxable": feuille.cell(row=ligne_donnees, column=12).value,
        "taxe_5": feuille.cell(row=ligne_donnees, column=13).value,
        "retenue_amicale": feuille.cell(row=ligne_donnees, column=14).value,
        "dette": feuille.cell(row=ligne_donnees, column=15).value,
        "net_a_percevoir": feuille.cell(row=ligne_donnees, column=16).value,
    }

    assert valeurs_excel["total_heures"] == resultat_moteur.total_heures == 100
    assert valeurs_excel["taux_horaire"] == resultat_moteur.taux_horaire == 2000
    assert valeurs_excel["gain_heures"] == resultat_moteur.gain_heures == 200000
    assert valeurs_excel["prime_ap_pp"] == resultat_moteur.prime_ap_pp == 20000
    assert valeurs_excel["surveillance_secretariat"] == resultat_moteur.surveillance_secretariat == 10000
    assert valeurs_excel["indemnite_suggestion_admin"] == resultat_moteur.indemnite_suggestion_admin == 5000
    assert valeurs_excel["base_taxable"] == resultat_moteur.base_taxable == 235000
    assert valeurs_excel["taxe_5"] == resultat_moteur.taxe_5 == 11750
    assert valeurs_excel["retenue_amicale"] == resultat_moteur.retenue_amicale == 5000
    assert valeurs_excel["dette"] == resultat_moteur.dette == 10000
    assert valeurs_excel["net_a_percevoir"] == resultat_moteur.net_a_percevoir == 208250
