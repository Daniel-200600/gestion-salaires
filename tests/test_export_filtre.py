"""
Tests de la correction n°1 : l'export Excel du tableau de bord doit
exporter EXACTEMENT les enseignants filtrés, jamais toute la période.

Chaque test relit le fichier généré (openpyxl) et vérifie les noms
réellement présents — il ne suffit jamais de vérifier que le fichier
existe.
"""

import pytest
from openpyxl import load_workbook

import database.connection as database_connection
from services import dashboard_service, donnees_paie_service, enseignant_service, periode_service
from services.comptabilite_service import ComptabiliteError
from services.donnees_paie_service import DonneesPaieEnseignant


@pytest.fixture(autouse=True)
def _rediriger_connexion_par_defaut(db_path, monkeypatch, tmp_path):
    monkeypatch.setattr(database_connection, "DB_PATH", db_path)
    monkeypatch.setattr(dashboard_service, "EXPORT_DIR", tmp_path / "exports")


def _creer_periode_avec_quatre_enseignants():
    """
    e1 : Permanent, Masculin, actif
    e2 : Vacataire, Féminin, actif
    e3 : Vacataire, Masculin, actif
    e4 : Permanent, Féminin, INACTIF
    """
    e1 = enseignant_service.creer_enseignant(nom="Ateba", prenom="Paul", sexe="M", statut="P", taux_horaire=1000)
    e2 = enseignant_service.creer_enseignant(nom="Biya", prenom="Aicha", sexe="F", statut="V", taux_horaire=1000)
    e3 = enseignant_service.creer_enseignant(nom="Chomdack", prenom="Karim", sexe="M", statut="V", taux_horaire=1000)
    e4 = enseignant_service.creer_enseignant(nom="Dupont", prenom="Alice", sexe="F", statut="P", taux_horaire=1000)

    p = periode_service.creer_periode(mois=5, annee=2035)
    p = periode_service.ouvrir_periode(p.id)
    for e in (e1, e2, e3, e4):
        donnees_paie_service.enregistrer_donnees_paie_groupe(
            p.id, [DonneesPaieEnseignant(enseignant_id=e.id, heures_par_semaine={1: 10})]
        )
    p = periode_service.valider_periode(p.id)
    enseignant_service.desactiver_enseignant(e4.id)
    return p, e1, e2, e3, e4


def _noms_dans_fichier(chemin):
    wb = load_workbook(chemin)
    ws = wb["État comptable"]
    return {c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)}


# ---------------------------------------------------------------------
# 1. Export sans filtre : tous les enseignants de la période
# ---------------------------------------------------------------------

def test_export_sans_filtre_contient_tous_les_enseignants():
    p, e1, e2, e3, e4 = _creer_periode_avec_quatre_enseignants()
    chemin = dashboard_service.exporter_consultation_excel(p.id, enseignant_ids=None)
    noms = _noms_dans_fichier(chemin)
    assert {"Ateba", "Biya", "Chomdack", "Dupont"}.issubset(noms)
    assert chemin.name == "Consultation_Paie_Mai_2035.xlsx"


# ---------------------------------------------------------------------
# 2. Export avec un seul enseignant
# ---------------------------------------------------------------------

def test_export_un_seul_enseignant():
    p, e1, e2, e3, e4 = _creer_periode_avec_quatre_enseignants()
    chemin = dashboard_service.exporter_consultation_excel(p.id, enseignant_ids=[e1.id])
    noms = _noms_dans_fichier(chemin)
    assert "Ateba" in noms
    assert "Biya" not in noms and "Chomdack" not in noms
    assert "ATEBA" in chemin.name.upper()


# ---------------------------------------------------------------------
# 3. Export avec plusieurs enseignants filtrés
# ---------------------------------------------------------------------

def test_export_plusieurs_enseignants_filtres():
    p, e1, e2, e3, e4 = _creer_periode_avec_quatre_enseignants()
    chemin = dashboard_service.exporter_consultation_excel(p.id, enseignant_ids=[e2.id, e3.id])
    noms = _noms_dans_fichier(chemin)
    assert {"Biya", "Chomdack"}.issubset(noms)
    assert "Ateba" not in noms
    assert chemin.name == "Consultation_Paie_Filtree_Mai_2035.xlsx"


# ---------------------------------------------------------------------
# 4. Export par statut (Vacataire)
# ---------------------------------------------------------------------

def test_export_filtre_par_statut():
    from models.enums import StatutEnseignant

    p, e1, e2, e3, e4 = _creer_periode_avec_quatre_enseignants()
    tous = enseignant_service.lister_enseignants(inclure_inactifs=True)
    vacataires = dashboard_service.filtrer_enseignants(tous, statut=StatutEnseignant.VACATAIRE)
    ids = [e.id for e in vacataires]

    chemin = dashboard_service.exporter_consultation_excel(p.id, enseignant_ids=ids)
    noms = _noms_dans_fichier(chemin)
    assert {"Biya", "Chomdack"}.issubset(noms)
    assert "Ateba" not in noms  # permanent exclu


# ---------------------------------------------------------------------
# 5. Export par sexe (Féminin)
# ---------------------------------------------------------------------

def test_export_filtre_par_sexe():
    from models.enums import Sexe

    p, e1, e2, e3, e4 = _creer_periode_avec_quatre_enseignants()
    tous = enseignant_service.lister_enseignants(inclure_inactifs=True)
    femmes = dashboard_service.filtrer_enseignants(tous, sexe=Sexe.FEMME)
    ids = [e.id for e in femmes]

    chemin = dashboard_service.exporter_consultation_excel(p.id, enseignant_ids=ids)
    noms = _noms_dans_fichier(chemin)
    assert "Biya" in noms
    assert "Ateba" not in noms and "Chomdack" not in noms


# ---------------------------------------------------------------------
# 6. Export par actif/inactif
# ---------------------------------------------------------------------

def test_export_filtre_actifs_uniquement():
    p, e1, e2, e3, e4 = _creer_periode_avec_quatre_enseignants()
    tous = enseignant_service.lister_enseignants(inclure_inactifs=True)
    actifs = dashboard_service.filtrer_enseignants(tous, actif=True)
    ids = [e.id for e in actifs]

    chemin = dashboard_service.exporter_consultation_excel(p.id, enseignant_ids=ids)
    noms = _noms_dans_fichier(chemin)
    assert {"Ateba", "Biya", "Chomdack"}.issubset(noms)
    assert "Dupont" not in noms  # inactif, exclu


# ---------------------------------------------------------------------
# 7. Export après recherche par nom
# ---------------------------------------------------------------------

def test_export_apres_recherche():
    p, e1, e2, e3, e4 = _creer_periode_avec_quatre_enseignants()
    resultats_recherche = enseignant_service.rechercher_enseignants("ateba", inclure_inactifs=True)
    ids = [e.id for e in resultats_recherche]

    chemin = dashboard_service.exporter_consultation_excel(p.id, enseignant_ids=ids)
    noms = _noms_dans_fichier(chemin)
    assert "Ateba" in noms
    assert "Biya" not in noms and "Chomdack" not in noms


# ---------------------------------------------------------------------
# 8. Export avec combinaison de filtres (statut + sexe)
# ---------------------------------------------------------------------

def test_export_combinaison_filtres():
    from models.enums import Sexe, StatutEnseignant

    p, e1, e2, e3, e4 = _creer_periode_avec_quatre_enseignants()
    tous = enseignant_service.lister_enseignants(inclure_inactifs=True)
    filtres = dashboard_service.filtrer_enseignants(tous, statut=StatutEnseignant.VACATAIRE, sexe=Sexe.HOMME)
    ids = [e.id for e in filtres]

    chemin = dashboard_service.exporter_consultation_excel(p.id, enseignant_ids=ids)
    noms = _noms_dans_fichier(chemin)
    assert "Chomdack" in noms  # vacataire + masculin
    assert "Ateba" not in noms  # permanent
    assert "Biya" not in noms  # féminin


# ---------------------------------------------------------------------
# 9. Aucun résultat : pas de fichier vide silencieux
# ---------------------------------------------------------------------

def test_export_aucun_resultat_leve_erreur_claire():
    p, e1, e2, e3, e4 = _creer_periode_avec_quatre_enseignants()
    with pytest.raises(ComptabiliteError, match="Aucune donnée"):
        dashboard_service.exporter_consultation_excel(p.id, enseignant_ids=[])


# ---------------------------------------------------------------------
# Vérification du nombre de lignes (pas seulement les noms)
# ---------------------------------------------------------------------

def test_export_nombre_de_lignes_correspond_exactement_au_filtre():
    p, e1, e2, e3, e4 = _creer_periode_avec_quatre_enseignants()
    chemin = dashboard_service.exporter_consultation_excel(p.id, enseignant_ids=[e1.id, e2.id])

    wb = load_workbook(chemin)
    ws = wb["État comptable"]
    lignes_avec_nom_attendu = sum(
        1 for row in ws.iter_rows(min_row=1) if row[1].value in ("Ateba", "Biya")
    )
    assert lignes_avec_nom_attendu == 2
