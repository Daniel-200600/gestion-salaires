"""
Import du fichier de paie mensuel (classeur Heures / Global / Comptable),
avec un classeur fictif construit sur le modèle de celui des établissements.
"""

import io
import sqlite3

import pytest
from openpyxl import Workbook

import database.connection as database_connection
from database.repositories import enseignant_repository
from models.enseignant import Enseignant
from models.enums import Sexe, StatutEnseignant
from services import enseignant_service, paie_service, periode_service
from services import import_fichier_paie_service as svc

ENTETES_GLOBAL = ["S/N", "Noms & Prenoms", "Sexe", "S1", "S2", "S3", "S4", "S5", "Total", "Status", "Taux Horaire",
                  "Gain Heures", "Prime AP/PP", "Surveillance/Secretariat", "Indemnite Suggestion/Admin", "Tax",
                  "Retenue Amicale", "Dette", "Net a percevoir"]


def _classeur(lignes, heures=None, comptable=None):
    """
    `lignes` : dicts (nom, sexe, statut, taux, gain, s=[...], prime, indemnite, retenue, net, formule_gain).
    La feuille des heures reprend `s` sauf indication contraire dans `heures` ({S/N: (nom, [s1..s5])}).
    """
    classeur = Workbook()
    feuille_h = classeur.active
    feuille_h.title = "Heures"
    feuille_h.append(["S/N", "Noms & Prenoms", "S1", "S2", "S3", "S4", "S5", "Total"])
    feuille_g = classeur.create_sheet("Global")
    feuille_g.append(ENTETES_GLOBAL)
    feuille_c = classeur.create_sheet("Comptable")
    feuille_c.append(["S/N", "Noms & Prenoms", "Net a percevoir"])
    for sn, l in enumerate(lignes, start=1):
        nom_h, s_h = (heures or {}).get(sn, (l["nom"], l["s"]))
        feuille_h.append([sn, nom_h, *s_h, None, sum(x or 0 for x in s_h)])
        ligne = sn + 1
        gain = f"=I{ligne}*K{ligne}" if l.get("formule_gain") else l.get("gain")
        feuille_g.append([sn, l["nom"], l["sexe"], *l["s"], None, sum(x or 0 for x in l["s"]), l["statut"],
                          l.get("taux"), gain, l.get("prime"), None, l.get("indemnite"), None,
                          l.get("retenue", 5000), None, l.get("net")])
        feuille_c.append([sn, l["nom"], (comptable or {}).get(sn, l.get("net"))])
    feuille_h.append([None, "Total"])
    feuille_g.append([None, "TOTAL SALAIRES"])
    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


VACATAIRE = {"nom": "BELLA NGONO Christine", "sexe": "F", "statut": "V", "taux": 1800, "formule_gain": True,
             "s": [10, 8, 0, 2], "net": 29020}  # 20 h × 1 800 = 36 000 ; taxe 1 980 ; amicale 5 000
PERMANENT_FIXE = {"nom": "ESSOMBA Jean", "sexe": "M", "statut": "P", "gain": 150000, "s": [2, 3, None, 3],
                  "indemnite": 5000, "net": 150000}
PERMANENT_HORAIRE = {"nom": "OWONA Luc", "sexe": "M", "statut": "P", "taux": 2000, "formule_gain": True,
                     "s": [5, 5, 5, 5], "net": 35000}  # 20 h × 2 000, sans taxe, amicale 5 000


@pytest.fixture(autouse=True)
def _base(db_path, monkeypatch):
    monkeypatch.setattr(database_connection, "DB_PATH", db_path)


@pytest.fixture
def periode():
    return periode_service.ouvrir_periode(periode_service.creer_periode(mois=9, annee=2026).id)


def _lire(*lignes, **options):
    return svc.lire_fichier_paie(_classeur(list(lignes), **options), "paie.xlsm")


def test_feuilles_reconnues_et_lignes_lues():
    lu = _lire(VACATAIRE, PERMANENT_FIXE, PERMANENT_HORAIRE)
    assert {f.nom: f.role for f in lu.feuilles} == {"Heures": "heures", "Global": "global", "Comptable": "comptable"}
    assert [l["nom_complet"] for l in lu.lignes] == ["BELLA NGONO Christine", "ESSOMBA Jean", "OWONA Luc"]
    assert lu.total_net_fichier == 29020 + 150000 + 35000
    vacataire, fixe, horaire = lu.lignes
    assert (vacataire["s1"], vacataire["s4"], vacataire["taux_horaire"], vacataire["salaire_fixe"]) == (10, 2, 1800, None)
    assert (fixe["salaire_fixe"], fixe["taux_horaire"], fixe["indemnite"]) == (150000, None, 5000)
    assert horaire["salaire_fixe"] is None and horaire["taux_horaire"] == 2000


def test_calcul_identique_au_fichier(periode):
    lignes = svc.verifier_lignes(_lire(VACATAIRE, PERMANENT_FIXE, PERMANENT_HORAIRE).lignes, periode.id)
    assert [(l.action, l.net_calcule, l.ecart, l.erreurs) for l in lignes] == [
        ("Création", 29020, 0, []), ("Création", 150000, 0, []), ("Création", 35000, 0, []),
    ]


def test_la_feuille_des_heures_fait_foi(periode):
    global_different = dict(VACATAIRE, s=[10, 9, 0, 2])  # S2 tapé à la main dans Global
    lu = _lire(global_different, heures={1: (VACATAIRE["nom"], [10, 8, 0, 2])})
    ligne = lu.lignes[0]
    assert ligne["s2"] == 8
    assert "S2 : 8 h dans la feuille des heures, 9 h dans la feuille Global" in ligne["remarques_fichier"]
    assert "Total du fichier : 21 h ; somme des semaines retenues : 20 h." in ligne["remarques_fichier"]


def test_nom_ecrit_autrement_et_ecart_de_net_signales(periode):
    lu = _lire(VACATAIRE, heures={1: ("BELLA Christine", VACATAIRE["s"])}, comptable={1: 29000})
    verifiee = svc.verifier_lignes(lu.lignes, periode.id)[0]
    assert "Nom écrit autrement dans la feuille des heures : « BELLA Christine »" in verifiee.remarques[0]
    assert verifiee.ecart == 20 and "écart de +20 FCFA" in verifiee.remarques[-1]


def test_erreurs_bloquantes_et_correction_dans_le_tableau(periode):
    lu = _lire(dict(VACATAIRE, taux=None, sexe=None), PERMANENT_FIXE)
    premiere = svc.verifier_lignes(lu.lignes, periode.id)[0]
    assert premiere.erreurs == ["Sexe manquant ou illisible (M ou F).", "Taux horaire manquant."]
    with pytest.raises(svc.ImportFichierPaieError, match="1 ligne"):
        svc.enregistrer_lignes(svc.verifier_lignes(lu.lignes, periode.id), periode.id, "paie.xlsm")

    lu.lignes[0].update(sexe="F", taux_horaire=1800)  # correction faite dans le tableau
    rapport = svc.enregistrer_lignes(svc.verifier_lignes(lu.lignes, periode.id), periode.id, "paie.xlsm")
    assert (rapport.nb_crees, rapport.nb_lignes) == (2, 2)


def test_enregistrement_complet_puis_calcul_de_la_periode(periode, db_path):
    lu = _lire(VACATAIRE, PERMANENT_FIXE, PERMANENT_HORAIRE)
    rapport = svc.enregistrer_lignes(svc.verifier_lignes(lu.lignes, periode.id), periode.id, "paie.xlsm",
                                     utilisateur="gestionnaire")
    assert (rapport.nb_crees, rapport.nb_mis_a_jour) == (3, 0)
    enseignants = {e.nom: e for e in enseignant_service.lister_enseignants()}
    assert (enseignants["BELLA NGONO"].prenom, enseignants["ESSOMBA"].salaire_fixe) == ("Christine", 150000)
    groupe = paie_service.calculer_paie_groupe(periode.id, [e.id for e in enseignants.values()])
    assert paie_service.calculer_totaux_groupe(groupe.resultats).total_net_a_percevoir == lu.total_net_fichier
    with sqlite3.connect(db_path) as conn:
        details = conn.execute("SELECT details FROM audit_log WHERE type_action = 'import_donnees'").fetchone()[0]
    assert "3 fiche(s) créée(s)" in details


def test_fiche_existante_mise_a_jour_sans_changer_son_nom(periode):
    eid = enseignant_repository.creer(Enseignant(nom="BELLA NGONO", prenom="Christine", sexe=Sexe.FEMME,
                                                 statut=StatutEnseignant.VACATAIRE, taux_horaire=1700,
                                                 telephone="600000000"))
    ligne = svc.verifier_lignes(_lire(dict(VACATAIRE, nom="Christine Bella Ngono")).lignes, periode.id)[0]
    assert ligne.action == "Mise à jour" and ligne.modifications == ["Taux horaire : 1700 → 1800"]
    svc.enregistrer_lignes([ligne], periode.id, "paie.xlsm")
    fiche = enseignant_service.obtenir_enseignant(eid)
    assert (fiche.nom, fiche.prenom, fiche.taux_horaire, fiche.telephone) == ("BELLA NGONO", "Christine", 1800, "600000000")


def test_ligne_decochee_ignoree_et_periode_non_ouverte_refusee(periode):
    lu = _lire(VACATAIRE, PERMANENT_FIXE)
    lu.lignes[1]["importer"] = False
    lignes = svc.verifier_lignes(lu.lignes, periode.id)
    assert [l.action for l in lignes] == ["Création", "Ignorée"]
    brouillon = periode_service.creer_periode(mois=10, annee=2026)
    with pytest.raises(svc.ImportFichierPaieError, match="ouverte"):
        svc.enregistrer_lignes(lignes, brouillon.id, "paie.xlsm")
    assert enseignant_service.lister_enseignants() == []  # rien d'écrit


def test_tout_ou_rien(periode, monkeypatch):
    from services import donnees_paie_service

    appels = []

    def ecrire(conn, periode_id, donnees):
        appels.append(donnees)
        if len(appels) == 2:
            raise ValueError("panne simulée")
    monkeypatch.setattr(donnees_paie_service, "ecrire_donnees_paie", ecrire)
    lignes = svc.verifier_lignes(_lire(VACATAIRE, PERMANENT_FIXE).lignes, periode.id)
    with pytest.raises(svc.ImportFichierPaieError, match="rien n'a été enregistré"):
        svc.enregistrer_lignes(lignes, periode.id, "paie.xlsm")
    assert enseignant_service.lister_enseignants() == []


def test_fichiers_refuses():
    with pytest.raises(svc.ImportFichierPaieError, match="Format"):
        svc.lire_fichier_paie(b"x", "paie.pdf")
    classeur = Workbook()
    classeur.active.append(["Nom", "Prénom"])
    tampon = io.BytesIO()
    classeur.save(tampon)
    with pytest.raises(svc.ImportFichierPaieError, match="Feuille des enseignants introuvable"):
        svc.lire_fichier_paie(tampon.getvalue(), "liste.xlsx")


@pytest.mark.mode_demo
def test_limite_du_mode_demonstration(periode, tmp_path, monkeypatch):
    from services import licence_service

    monkeypatch.setattr(licence_service, "chemin_licence", lambda: tmp_path / "licence.cle")
    lignes = [dict(VACATAIRE, nom=f"{nom} Paul") for nom in ("ATEBA", "BIYA", "ESSO", "FOUDA", "MBIDA", "NANA")]
    verifiees = svc.verifier_lignes(_lire(*lignes).lignes, periode.id)
    with pytest.raises(svc.ImportFichierPaieError, match="Mode démonstration"):
        svc.enregistrer_lignes(verifiees, periode.id, "paie.xlsm")


def test_valeurs_decimales_du_tableau_modifie(periode):
    # Le tableau de la page rend les nombres en décimaux (1800.0) et les cases vides en NaN.
    ligne = dict(_lire(VACATAIRE).lignes[0], taux_horaire=1800.0, net_fichier=29020.0, sn=1.0, dette=float("nan"))
    verifiee = svc.verifier_lignes([ligne], periode.id)[0]
    assert (verifiee.erreurs, verifiee.ecart, verifiee.donnees["sn"], verifiee.donnees["dette"]) == ([], 0, 1, None)


def test_nom_absent_ne_bloque_pas(periode):
    lu = _lire(dict(VACATAIRE, nom="SANS-IDENTITE"))
    lu.lignes[0]["nom_complet"] = None  # case du nom vidée dans le tableau
    ligne = svc.verifier_lignes(lu.lignes, periode.id)[0]
    assert ligne.erreurs == [] and ligne.donnees["nom_complet"] == "SANS NOM 1"
    assert "nom provisoire « SANS NOM 1 »" in ligne.remarques[0]
    svc.enregistrer_lignes([ligne], periode.id, "paie.xlsm")
    assert enseignant_service.lister_enseignants()[0].nom == "SANS NOM 1"


def test_nom_d_un_seul_mot_accepte_jusqu_au_bulletin(periode):
    from services import bulletin_service, controle_paie_service

    ligne = svc.verifier_lignes(_lire(dict(VACATAIRE, nom="TCHOUA")).lignes, periode.id)[0]
    assert ligne.erreurs == []
    svc.enregistrer_lignes([ligne], periode.id, "paie.xlsm")
    enseignant = enseignant_service.lister_enseignants()[0]
    assert (enseignant.nom, enseignant.prenom) == ("TCHOUA", "")
    periode = periode_service.valider_periode(periode.id)  # aucune erreur bloquante
    assert not controle_paie_service.controler_periode(periode.id).est_bloque
    bulletin = bulletin_service.generer_bulletin_enseignant(periode.id, enseignant.id)
    assert bulletin.chemin.exists()


def test_taux_nul_et_net_negatif_bloquent_des_l_import(periode):
    taux_nul = dict(VACATAIRE, nom="ATEBA Paul", taux=0)
    sans_heures = dict(VACATAIRE, nom="BIYA Rene", s=[None, None, None, None])  # 5 000 de retenue, aucun gain
    lignes = svc.verifier_lignes(_lire(taux_nul, sans_heures).lignes, periode.id)
    assert lignes[0].erreurs == ["Taux horaire nul : indiquez le taux (ou un salaire fixe pour un permanent)."]
    assert lignes[1].net_calcule == -5000 and lignes[1].erreurs[0].startswith("Net négatif (-5000 FCFA)")


def test_bulletin_d_un_nom_compose(periode):
    from services import bulletin_service

    svc.enregistrer_lignes(svc.verifier_lignes(_lire(VACATAIRE).lignes, periode.id), periode.id, "paie.xlsm")
    enseignant = enseignant_service.lister_enseignants()[0]
    assert enseignant.nom == "BELLA NGONO"  # nom composé : « BELLA_NGONO » dans le nom du fichier
    assert bulletin_service.generer_bulletin_enseignant(periode.id, enseignant.id).chemin.exists()


def test_rapprochement_avec_les_fiches_de_l_application(periode):
    # Fiches saisies avant l'import : nom complet dans « Nom », prénom vide ; nom avec un prénom en plus.
    sans_prenom = enseignant_repository.creer(Enseignant(nom="OKALA DIDIER", sexe=Sexe.HOMME,
                                                         statut=StatutEnseignant.VACATAIRE, taux_horaire=1700))
    prenom_en_plus = enseignant_repository.creer(Enseignant(nom="TABI", prenom="NGWA BERNARD", sexe=Sexe.HOMME,
                                                            statut=StatutEnseignant.VACATAIRE, taux_horaire=1800))
    lu = _lire(dict(VACATAIRE, nom="OKALA DIDIER", sexe="M", taux=1700),
               dict(VACATAIRE, nom="TABI BERNARD", sexe="M", taux=1800),
               dict(VACATAIRE, nom="NOUVEAU Venu"))
    lignes = svc.verifier_lignes(lu.lignes, periode.id)
    assert [l.action for l in lignes] == ["Mise à jour", "Mise à jour", "Création"]
    assert lignes[0].modifications == ["Nom et prénom : « OKALA DIDIER » → nom « OKALA », prénom « DIDIER »"]
    assert "Rapproché de la fiche « TABI NGWA BERNARD »" in lignes[1].remarques[0]
    svc.enregistrer_lignes(lignes, periode.id, "paie.xlsm")
    assert len(enseignant_service.lister_enseignants()) == 3  # aucun doublon créé
    fiche = enseignant_service.obtenir_enseignant(sans_prenom)
    assert (fiche.nom, fiche.prenom) == ("OKALA", "DIDIER")
    assert enseignant_service.obtenir_enseignant(prenom_en_plus).prenom == "NGWA BERNARD"  # nom complet conservé


def test_nom_proche_de_plusieurs_fiches_cree_une_nouvelle(periode):
    for prenom in ("Marie Claire", "Marie Louise"):
        enseignant_repository.creer(Enseignant(nom="NGONO", prenom=prenom, sexe=Sexe.FEMME,
                                               statut=StatutEnseignant.VACATAIRE, taux_horaire=1800))
    ligne = svc.verifier_lignes(_lire(dict(VACATAIRE, nom="NGONO Marie")).lignes, periode.id)[0]
    assert ligne.action == "Création" and "Plusieurs fiches ont un nom proche" in ligne.remarques[0]
