"""
Import d'une liste d'enseignants existante (Excel, Word ou PDF), avec des
fiches incomplètes à compléter plus tard.
"""

import pandas as pd
import pytest
from docx import Document

from models.enums import Sexe, StatutEnseignant, StrategieDoublon, TypeImport
from services import enseignant_service, import_service, paie_service, periode_service
from services.donnees_paie_service import DonneesPaieEnseignant, enregistrer_donnees_paie_groupe
from services.import_service import ImportServiceError
from services.paie_service import CalculPaieError
from utils.validators import EnseignantValidationError

LISTE = [
    ["Nom et prénoms", "Genre", "Catégorie", "Taux/heure (FCFA)", "Contact"],
    ["MBARGA Élise", "Féminin", "Vacataire", "1 500", "699 00 00 01"],
    ["ATEBA Paul", "Masculin", "", "", ""],
    ["NGONO Marie Claire", "", "Permanent", "2.000 F", ""],
]


def _word(tmp_path, lignes=LISTE, nom="liste.docx"):
    document = Document()
    document.add_paragraph("LISTE DES ENSEIGNANTS — ANNÉE 2026-2027")
    tableau = document.add_table(rows=len(lignes), cols=len(lignes[0]))
    for i, ligne in enumerate(lignes):
        for j, valeur in enumerate(ligne):
            tableau.cell(i, j).text = valeur
    chemin = tmp_path / nom
    document.save(chemin)
    return chemin


def _pdf(tmp_path, pages, nom="liste.pdf"):
    """PDF à tableau quadrillé ; `pages` = liste de listes de lignes (une par page)."""
    import pymupdf

    document = pymupdf.open()
    for lignes in pages:
        page = document.new_page()
        largeur, hauteur, x0, y0 = 105, 24, 40, 60
        for i, ligne in enumerate(lignes):
            for j, valeur in enumerate(ligne):
                rect = pymupdf.Rect(x0 + j * largeur, y0 + i * hauteur, x0 + (j + 1) * largeur, y0 + (i + 1) * hauteur)
                page.draw_rect(rect, color=(0, 0, 0), width=0.8)
                if valeur:
                    page.insert_text((rect.x0 + 3, rect.y0 + 15), valeur, fontsize=8)
    chemin = tmp_path / nom
    document.save(chemin)
    return chemin


def _preparer(chemin, nom_fichier, strategie=StrategieDoublon.REFUSER, db_path=None):
    contenu = import_service.lire_fichier(chemin, nom_fichier)
    analyse = import_service.analyser_fichier(contenu, TypeImport.ENSEIGNANTS)
    return analyse, import_service.preparer_import_enseignants(
        analyse, contenu.feuilles[analyse.feuille_choisie], strategie, db_path=db_path
    )


def _importer(chemin, nom_fichier, db_path, strategie=StrategieDoublon.REFUSER):
    _, rapport = _preparer(chemin, nom_fichier, strategie, db_path=db_path)
    import_service.executer_import_enseignants(rapport, nom_fichier, db_path=db_path)
    return rapport


def _par_nom(db_path):
    return {e.nom: e for e in enseignant_service.lister_enseignants(db_path=db_path)}


# ---------------------------------------------------------------------
# Lecture Word et PDF
# ---------------------------------------------------------------------

def test_liste_word_lue_et_colonnes_reconnues(tmp_path, db_path):
    analyse, rapport = _preparer(_word(tmp_path), "liste.docx", db_path=db_path)
    assert analyse.feuille_choisie == "Tableau 1" and analyse.nombre_lignes == 3
    assert set(analyse.colonnes_reconnues.values()) == {"nom_complet", "sexe", "statut", "taux_horaire", "telephone"}
    assert rapport.nb_creations == 3 and rapport.nb_erreurs == 0


def test_liste_pdf_sur_deux_pages_reunie(tmp_path, db_path):
    chemin = _pdf(tmp_path, [LISTE[:3], LISTE[3:]])  # la 2e page continue le tableau, sans en-tête
    contenu = import_service.lire_fichier(chemin, "liste.pdf")
    assert list(contenu.feuilles) == ["Tableau 1"]
    _, rapport = _preparer(chemin, "liste.pdf", db_path=db_path)
    assert [l.donnees["nom"] for l in rapport.lignes] == ["MBARGA", "ATEBA", "NGONO"]


def test_word_sans_tableau_refuse_avec_explication(tmp_path):
    document = Document()
    document.add_paragraph("MBARGA Élise")
    document.save(tmp_path / "texte.docx")
    with pytest.raises(ImportServiceError, match="tableau"):
        import_service.lire_fichier(tmp_path / "texte.docx", "texte.docx")


def test_pdf_scanne_refuse_avec_explication(tmp_path):
    import pymupdf

    document = pymupdf.open()
    document.new_page().draw_rect(pymupdf.Rect(50, 50, 200, 200))  # aucun texte
    document.save(tmp_path / "scan.pdf")
    with pytest.raises(ImportServiceError, match="scanné"):
        import_service.lire_fichier(tmp_path / "scan.pdf", "scan.pdf")


# ---------------------------------------------------------------------
# Fiches incomplètes
# ---------------------------------------------------------------------

def test_valeurs_interpretees_et_fiches_incompletes_creees(tmp_path, db_path):
    rapport = _importer(_word(tmp_path), "liste.docx", db_path)
    assert rapport.nb_fiches_incompletes == 2

    enseignants = _par_nom(db_path)
    mbarga = enseignants["MBARGA"]
    assert (mbarga.prenom, mbarga.sexe, mbarga.statut, mbarga.taux_horaire) == (
        "Élise", Sexe.FEMME, StatutEnseignant.VACATAIRE, 1500)
    assert mbarga.telephone == "699 00 00 01" and mbarga.est_complet
    assert enseignants["ATEBA"].champs_manquants == ["statut", "taux horaire"]
    assert enseignants["NGONO"].prenom == "Marie Claire"
    assert enseignants["NGONO"].champs_manquants == ["sexe"] and enseignants["NGONO"].taux_horaire == 2000


def test_seul_le_nom_est_obligatoire(tmp_path, db_path):
    df = pd.DataFrame({"Nom": ["", "FOUDA"], "Prénom": ["Alain", ""], "Sexe": ["M", ""]})
    df.to_excel(tmp_path / "l.xlsx", index=False)
    _, rapport = _preparer(tmp_path / "l.xlsx", "l.xlsx", db_path=db_path)
    assert rapport.nb_rejetees == 1 and rapport.nb_creations == 1
    assert rapport.lignes[1].donnees["prenom"] == ""


def test_reimport_complete_sans_rien_effacer(tmp_path, db_path):
    _importer(_word(tmp_path), "liste.docx", db_path)
    complement = [["Nom", "Prénom", "Statut", "Taux horaire", "Contact"],
                  ["ATEBA", "Paul", "Permanent", "1800", ""],
                  ["MBARGA", "Élise", "", "", ""]]  # rien de nouveau : ne doit rien effacer
    _importer(_word(tmp_path, complement, "complement.docx"), "complement.docx", db_path,
              strategie=StrategieDoublon.METTRE_A_JOUR)

    enseignants = _par_nom(db_path)
    assert enseignants["ATEBA"].statut == StatutEnseignant.PERMANENT and enseignants["ATEBA"].taux_horaire == 1800
    assert enseignants["ATEBA"].sexe == Sexe.HOMME  # déjà connu, conservé
    assert enseignants["MBARGA"].taux_horaire == 1500 and enseignants["MBARGA"].telephone == "699 00 00 01"


# ---------------------------------------------------------------------
# Paie : une fiche incomplète n'y entre pas tant qu'elle n'est pas complétée
# ---------------------------------------------------------------------

def _periode_ouverte(db_path):
    periode = periode_service.creer_periode(mois=4, annee=2030, db_path=db_path)
    return periode_service.ouvrir_periode(periode.id, db_path=db_path)


def test_fiche_incomplete_exclue_de_la_paie_puis_completee(tmp_path, db_path):
    _importer(_word(tmp_path), "liste.docx", db_path)
    ateba = _par_nom(db_path)["ATEBA"]
    periode = _periode_ouverte(db_path)

    assert ateba.id not in [e.id for e in enseignant_service.lister_enseignants_payables(db_path=db_path)]
    assert ateba.id in [e.id for e in enseignant_service.lister_enseignants_a_completer(db_path=db_path)]
    with pytest.raises(Exception, match="incomplète"):
        enregistrer_donnees_paie_groupe(
            periode.id, [DonneesPaieEnseignant(enseignant_id=ateba.id, heures_par_semaine={1: 4})], db_path=db_path
        )
    with pytest.raises(CalculPaieError, match="incomplète"):
        paie_service.calculer_paie_enseignant(periode.id, ateba.id, db_path=db_path)

    # Complétée depuis la page Enseignants : le statut et le taux suffisent,
    # le sexe déjà connu n'est pas à ressaisir.
    enseignant_service.modifier_enseignant(ateba.id, "ATEBA", "Paul", None, "P", 1800, db_path=db_path)
    assert enseignant_service.obtenir_enseignant(ateba.id, db_path=db_path).est_complet
    enregistrer_donnees_paie_groupe(
        periode.id, [DonneesPaieEnseignant(enseignant_id=ateba.id, heures_par_semaine={1: 4})], db_path=db_path
    )
    assert paie_service.calculer_paie_enseignant(periode.id, ateba.id, db_path=db_path).gain_heures == 7200


def test_completer_n_efface_jamais_une_valeur_et_controle_les_saisies(tmp_path, db_path):
    _importer(_word(tmp_path), "liste.docx", db_path)
    ngono = _par_nom(db_path)["NGONO"]
    enseignant_service.modifier_enseignant(ngono.id, "NGONO", "Marie Claire", "F", None, None, db_path=db_path)
    ngono = enseignant_service.obtenir_enseignant(ngono.id, db_path=db_path)
    assert (ngono.sexe, ngono.statut, ngono.taux_horaire) == (Sexe.FEMME, StatutEnseignant.PERMANENT, 2000)
    with pytest.raises(EnseignantValidationError):
        enseignant_service.modifier_enseignant(ngono.id, "NGONO", "Marie Claire", "Z", None, None, db_path=db_path)


def test_import_d_heures_refuse_pour_une_fiche_incomplete(tmp_path, db_path):
    _importer(_word(tmp_path), "liste.docx", db_path)
    periode = _periode_ouverte(db_path)
    pd.DataFrame({"Nom": ["ATEBA", "MBARGA"], "Prénom": ["Paul", "Élise"], "Semaine 1": [4, 6]}).to_excel(
        tmp_path / "heures.xlsx", index=False)
    contenu = import_service.lire_fichier(tmp_path / "heures.xlsx", "heures.xlsx")
    analyse = import_service.analyser_fichier(contenu, TypeImport.HEURES)
    rapport = import_service.preparer_import_heures(analyse, contenu.feuilles[analyse.feuille_choisie], periode.id,
                                                    db_path=db_path)
    erreurs = [a.message for a in rapport.toutes_anomalies if "incomplète" in a.message]
    assert len(erreurs) == 1 and "ATEBA" in erreurs[0]


def test_alerte_groupee_fiches_incompletes_puis_resolue(tmp_path, db_path):
    from services import alert_detection_service
    from database.repositories import alerte_repository

    _importer(_word(tmp_path), "liste.docx", db_path)
    alert_detection_service.detecter_alertes_enseignants(db_path=db_path)
    actives = [a for a in alerte_repository.lister_actives_par_source(
        alert_detection_service.SOURCE_ENSEIGNANTS, db_path=db_path) if a.type_alerte == "FICHES_INCOMPLETES"]
    assert len(actives) == 1  # une seule alerte, pas une par enseignant
    assert actives[0].titre.startswith("2 fiche(s)")

    for enseignant in enseignant_service.lister_enseignants_a_completer(db_path=db_path):
        enseignant_service.modifier_enseignant(
            enseignant.id, enseignant.nom, enseignant.prenom, sexe="M", statut="V", taux_horaire=1000,
            db_path=db_path,
        )
    alert_detection_service.detecter_alertes_enseignants(db_path=db_path)
    assert not [a for a in alerte_repository.lister_actives_par_source(
        alert_detection_service.SOURCE_ENSEIGNANTS, db_path=db_path) if a.type_alerte == "FICHES_INCOMPLETES"]
