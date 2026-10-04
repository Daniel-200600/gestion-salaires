"""Logo de l'établissement sur chaque classeur Excel exporté."""

import io

from openpyxl import Workbook, load_workbook
from PIL import Image

from exports import logo_excel
from exports.import_template_export import generer_classeur_modele
from models.enums import TypeImport


def _png(largeur=200, hauteur=100) -> bytes:
    tampon = io.BytesIO()
    Image.new("RGB", (largeur, hauteur), (31, 58, 110)).save(tampon, format="PNG")
    return tampon.getvalue()


def test_logo_sur_chaque_feuille_sans_deplacer_les_cellules():
    classeur = Workbook()
    classeur.active.append(["Nom", "Net"])
    classeur.create_sheet("Synthèse").append(["Total", 10])
    logo_excel.ajouter_logo(classeur, logo=_png())
    assert all(len(feuille._images) == 1 for feuille in classeur.worksheets)
    image = classeur.active._images[0]
    assert image.height == logo_excel.HAUTEUR_LOGO_PX and image.width == 128  # proportions gardées
    tampon = io.BytesIO()
    classeur.save(tampon)
    relu = load_workbook(io.BytesIO(tampon.getvalue()))
    assert [c.value for c in relu.active[1]] == ["Nom", "Net"]


def test_sans_logo_ou_logo_illisible_le_classeur_est_inchange(monkeypatch):
    monkeypatch.setattr(logo_excel, "logo_etablissement", lambda: None)
    classeur = logo_excel.ajouter_logo(Workbook())
    assert classeur.active._images == []
    assert logo_excel.ajouter_logo(Workbook(), logo=b"pas une image").active._images == []


def test_les_exports_portent_le_logo(monkeypatch):
    monkeypatch.setattr(logo_excel, "logo_etablissement", _png)
    classeur = generer_classeur_modele(TypeImport.ENSEIGNANTS)
    assert all(len(feuille._images) == 1 for feuille in classeur.worksheets)
