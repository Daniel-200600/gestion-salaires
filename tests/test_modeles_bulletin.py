"""
Bulletins de solde : reproduction du bulletin officiel et modèles importés (Word, PDF).

Couvre :
- le modèle Word standard (structure du bulletin officiel : logo, bandes vertes et
  jaunes, colonnes, mois en anglais, zone « Fait à ... le : » vierge) ;
- le modèle PDF standard (zones détectées automatiquement, fond conservé) ;
- l'import d'un modèle à balises et d'un bulletin déjà rempli, en Word et en PDF ;
- les refus (format, contenu, balises inconnues ou obligatoires manquantes, PDF sans texte) ;
- l'activation, la suppression, le repli sur le modèle standard ;
- la génération réelle des bulletins avec le modèle actif ;
- la réservation à l'administrateur (vérifiée côté service) et l'audit ;
- la conservation des modèles lors d'une réinitialisation des données.
"""

import ast
import io
import sqlite3
from pathlib import Path

import pymupdf
import pytest
from docx import Document

import database.connection as database_connection
from exports import pdf_export, word_export
from models.enums import RoleUtilisateur, TypeActionAudit
from services import (
    administration_service,
    auth_service,
    bulletin_service,
    donnees_paie_service,
    enseignant_service,
    modele_bulletin_service,
    periode_service,
    utilisateur_service,
)
from services.autorisation_service import AutorisationRefuseeError
from services.donnees_paie_service import DonneesPaieEnseignant
from services.modele_bulletin_service import ModeleBulletinError
from templates import build_template
from utils.balises_bulletin import BALISES, valeurs_exemple
from utils.detection_bulletin import suggerer_correspondances

# Cas de référence officiel : vacataire, taxe de 5 %.
pytestmark = pytest.mark.taxe_historique

RACINE = Path(__file__).resolve().parent.parent
MDP = "Motdepasse-Solide-1"
STANDARD = set(build_template.BALISES_MODELE_PDF_STANDARD)
VERT = (0, 176, 80)


@pytest.fixture(autouse=True)
def _environnement(db_path, monkeypatch, tmp_path):
    monkeypatch.setattr(database_connection, "DB_PATH", db_path)
    monkeypatch.setattr(auth_service, "_tentatives_echouees", {})
    monkeypatch.setattr(bulletin_service, "EXPORT_DIR_BULLETINS", tmp_path / "bulletins")
    monkeypatch.setattr(modele_bulletin_service, "MODELES_BULLETIN_DIR", tmp_path / "modeles")


# ---------------------------------------------------------------------
# Outils
# ---------------------------------------------------------------------

def _docx(document) -> bytes:
    return word_export.document_en_octets(document)


def _docx_a_balises(balises=("NOM_COMPLET", "NET_A_PAYER", "PERIODE", "STATUT", "TAXE")) -> bytes:
    document = Document()
    document.add_paragraph("BULLETIN DE PAIE — ÉTABLISSEMENT")
    table = document.add_table(rows=len(balises), cols=2)
    for index, balise in enumerate(balises):
        table.cell(index, 0).text = balise.title()
        table.cell(index, 1).text = "{{" + balise + "}}"
    return _docx(document)


def _pdf_a_balises() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.draw_rect(pymupdf.Rect(40, 60, 560, 90), color=None, fill=(0, 176 / 255, 80 / 255))
    page.insert_text((60, 80), "BULLETIN", fontname="Times-Bold", fontsize=11)
    page.insert_text((300, 80), "{{PERIODE|centre}}", fontname="Times-Bold", fontsize=11)
    page.insert_text((60, 120), "{{NOM_COMPLET}}", fontname="Times-Bold", fontsize=11)
    page.insert_text((400, 120), "Statut: {{STATUT}}", fontname="Times-Bold", fontsize=11)
    page.insert_text((60, 160), "NET A PAYER", fontname="Helvetica-Bold", fontsize=10)
    page.insert_text((400, 160), "{{NET_A_PAYER|droite}}", fontname="Helvetica-Bold", fontsize=10)
    return document.tobytes()


def _pdf_standard() -> bytes:
    return modele_bulletin_service.modele_standard_pdf().contenu()


def _docx_rempli() -> bytes:
    return _docx(build_template.construire_document("exemple"))


def _periode_validee_avec_enseignant(statut="V", taux_horaire=1800, heures=10, mois=7, annee=2026, **autres):
    e = enseignant_service.creer_enseignant(nom="Mbarga", prenom="Élise", sexe="F", statut=statut,
                                            taux_horaire=taux_horaire)
    p = periode_service.ouvrir_periode(periode_service.creer_periode(mois=mois, annee=annee).id)
    donnees_paie_service.enregistrer_donnees_paie_groupe(
        p.id, [DonneesPaieEnseignant(enseignant_id=e.id, heures_par_semaine={1: heures}, **autres)]
    )
    return e, periode_service.valider_periode(p.id)


def _texte_docx(chemin) -> str:
    return word_export.texte_document(Document(chemin))


def _couleur_pixel(contenu_pdf: bytes, x: float, y: float):
    document = pymupdf.open(stream=contenu_pdf, filetype="pdf")
    pixmap = document[0].get_pixmap()
    couleur = pixmap.pixel(int(x), int(y))
    document.close()
    return couleur


# ---------------------------------------------------------------------
# Modèle Word standard : structure du bulletin officiel
# ---------------------------------------------------------------------

def test_modele_word_standard_reproduit_la_structure_officielle():
    document = Document(word_export.TEMPLATE_PATH)
    texte = word_export.texte_document(document)
    for fixe in ("REPUBLIQUE DU CAMEROUN", "REPUBLIC OF CAMEROON",
                 "BULLETIN DE SOLDE / PAYSLIP", "ELEMENTS DE RENUMERATION / SALARY RUBRICS", "MONTANT / AMOUNT",
                 "S/N", "DESIGNATION", "GAINS", "RETENUES", "Gain Heures / Hourly Wage",
                 "Prime AP/PP / Incentive HOD/CM", "Surveillance/Secretariat / Invigilation",
                 "Indemnite Suggestion / Duty Post Allowance", "Taxe / Tax", "Retenue Amicale / Social Deduction",
                 "Dette / Debt", "Total", "NET A PAYER", "Done at Yaoundé on the / Fait à Yaoundé le:",
                 "Le Coordonateur Général/The General Coordinator", "Statut: "):
        assert fixe in texte, fixe
    xml = document.element.xml
    assert 'w:fill="00B050"' in xml and 'w:fill="FFFF00"' in xml  # bandes vertes et jaunes
    # Logo : présent seulement dans la version propre à l'établissement
    # (build_template.py --etablissement), jamais dans le dépôt public.
    assert len(document.inline_shapes) <= 1


def test_modele_word_standard_date_et_signature_vierges():
    """Comme sur le bulletin officiel : date et signature manuscrites."""
    texte = word_export.texte_document(Document(word_export.TEMPLATE_PATH))
    assert "{{DATE" not in texte
    assert texte.count("Fait à Yaoundé le:") == 1


def test_bulletin_genere_en_word_avec_mois_en_anglais_et_montants_bruts():
    e, p = _periode_validee_avec_enseignant()
    resultat = bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    assert resultat.chemin.suffix == ".docx"
    texte = _texte_docx(resultat.chemin)
    assert "JULY 2026" in texte
    assert "MBARGA" in texte and "ÉLISE" in texte and "Statut: V" in texte
    assert "18 000" in texte and "900" in texte  # 5 %, montants avec séparateur de milliers (version 1.8.0)
    assert "Seventeen Thousand One Hundred" in texte  # 17 100


# ---------------------------------------------------------------------
# Modèle PDF standard
# ---------------------------------------------------------------------

def test_modele_pdf_standard_zones_detectees_automatiquement():
    modele = modele_bulletin_service.modele_standard_pdf()
    assert {b for z in modele.zones for b in z.balises()} == STANDARD
    segments = pdf_export.segments_pdf(modele.contenu())
    assert set(suggerer_correspondances(segments).values()) == STANDARD


def test_modele_pdf_standard_ne_contient_aucune_donnee_personnelle():
    texte = pdf_export.texte_pdf(_pdf_standard())
    assert "NOM PRENOMS" in texte


def test_bulletin_pdf_avec_modele_standard(tmp_path):
    modele_bulletin_service.activer_modele(modele_bulletin_service.CLE_STANDARD_PDF)
    e, p = _periode_validee_avec_enseignant(taux_horaire=2000, heures=100, prime_ap_pp=35000, retenue_amicale=15000)
    resultat = bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    assert resultat.chemin.suffix == ".pdf"
    texte = pdf_export.texte_pdf(resultat.chemin.read_bytes())
    assert "Statut:" in texte and "\nV\n" in texte  # statut réécrit à droite de « Statut: »
    for valeur in ("MBARGA ÉLISE", "JULY 2026", "200 000", "35 000", "11 750", "15 000", "235 000",
                   "26 750", "208 250", "Two Hundred Eight Thousand Two Hundred Fifty"):
        assert valeur in texte, valeur
    assert "NOM PRENOMS" not in texte and "17010" not in texte  # valeurs d'exemple effacées
    assert "{{" not in texte
    assert resultat.chemin in bulletin_service.lister_bulletins_existants(p.libelle)


def test_fond_colore_conserve_sous_les_valeurs_remplacees():
    modele = modele_bulletin_service.modele_standard_pdf()
    zone = next(z for z in modele.zones if z.balises() == ["TOTAL_GAINS"])
    rendu = modele_bulletin_service.produire("pdf", modele.contenu(), modele.zones,
                                             {**valeurs_exemple(), "TOTAL_GAINS": "1"})
    x0, y0, x1, y1 = zone.rect
    assert _couleur_pixel(rendu, x1 - 1, (y0 + y1) / 2) == VERT


# ---------------------------------------------------------------------
# Import : modèles à balises
# ---------------------------------------------------------------------

def test_import_word_a_balises_et_generation():
    modele = modele_bulletin_service.importer_modele("Mon modèle", _docx_a_balises(), "modele.docx", activer=True)
    assert modele.chemin.exists() and modele.format == "docx"
    e, p = _periode_validee_avec_enseignant()
    texte = _texte_docx(bulletin_service.generer_bulletin_enseignant(p.id, e.id).chemin)
    assert "MBARGA ÉLISE" in texte and "17 100" in texte and "JULY 2026" in texte and "900" in texte


def test_import_pdf_a_balises_alignements_et_texte_fixe():
    analyse = modele_bulletin_service.analyser_modele(_pdf_a_balises(), "modele.pdf")
    assert analyse.mode == "balises" and analyse.valide
    gabarits = {z.gabarit: z.alignement for z in analyse.zones}
    assert gabarits == {"{PERIODE}": "centre", "{NOM_COMPLET}": "gauche", "Statut: {STATUT}": "gauche",
                        "{NET_A_PAYER}": "droite"}
    modele_bulletin_service.importer_modele("PDF balisé", _pdf_a_balises(), "modele.pdf", activer=True)
    e, p = _periode_validee_avec_enseignant(statut="P")
    texte = pdf_export.texte_pdf(bulletin_service.generer_bulletin_enseignant(p.id, e.id).chemin.read_bytes())
    # Permanent : aucune taxe, net = 10 h × 1 800 FCFA.
    assert "Statut: P" in texte and "MBARGA ÉLISE" in texte and "18 000" in texte and "{{" not in texte


def test_balises_historiques_acceptees():
    contenu = _docx_a_balises(("NOM", "PRENOM", "NET_A_PERÇEVOIR", "TAXE_5", "DATE_GENERATION"))
    analyse = modele_bulletin_service.analyser_modele(contenu, "ancien.docx")
    assert analyse.valide


def test_balise_inconnue_refusee():
    analyse = modele_bulletin_service.analyser_modele(_docx_a_balises(("NOM", "NET_A_PAYER", "SALAIRE")), "m.docx")
    assert not analyse.valide and analyse.inconnues == ["SALAIRE"]
    with pytest.raises(ModeleBulletinError, match="inconnue"):
        modele_bulletin_service.importer_modele("X", _docx_a_balises(("NOM", "NET_A_PAYER", "SALAIRE")), "m.docx")


def test_balises_obligatoires_manquantes_refusees():
    with pytest.raises(ModeleBulletinError, match="obligatoire"):
        modele_bulletin_service.importer_modele("X", _docx_a_balises(("PERIODE", "TAXE")), "m.docx")


# ---------------------------------------------------------------------
# Import : bulletins déjà remplis (correspondances proposées puis validées)
# ---------------------------------------------------------------------

def test_bulletin_word_rempli_detecte_et_converti():
    analyse = modele_bulletin_service.analyser_modele(_docx_rempli(), "bulletin.docx")
    assert analyse.mode == "correspondances"
    assert set(analyse.propositions.values()) == STANDARD
    modele = modele_bulletin_service.importer_modele(
        "Bulletin Word rempli", _docx_rempli(), "bulletin.docx", correspondances=analyse.propositions, activer=True
    )
    balises = word_export.extraire_balises_document(Document(modele.chemin))
    assert set(balises) == STANDARD
    e, p = _periode_validee_avec_enseignant()
    texte = _texte_docx(bulletin_service.generer_bulletin_enseignant(p.id, e.id).chemin)
    assert "MBARGA ÉLISE" in texte and "Statut: V" in texte and "17 100" in texte and "NOM PRENOMS" not in texte


def test_bulletin_pdf_rempli_detecte_et_reproduit():
    analyse = modele_bulletin_service.analyser_modele(_pdf_standard(), "bulletin.pdf")
    assert analyse.mode == "correspondances"
    assert set(analyse.propositions.values()) == STANDARD
    modele_bulletin_service.importer_modele(
        "Bulletin PDF rempli", _pdf_standard(), "bulletin.pdf", correspondances=analyse.propositions, activer=True
    )
    e, p = _periode_validee_avec_enseignant()
    texte = pdf_export.texte_pdf(bulletin_service.generer_bulletin_enseignant(p.id, e.id).chemin.read_bytes())
    assert "MBARGA ÉLISE" in texte and "17 100" in texte and "Seventeen Thousand One Hundred" in texte


def test_modele_pdf_enregistre_sans_les_valeurs_d_origine():
    """Les valeurs du bulletin envoyé (nom, montants) ne sont pas conservées dans le modèle enregistré."""
    analyse = modele_bulletin_service.analyser_modele(_pdf_standard(), "bulletin.pdf")
    modele = modele_bulletin_service.importer_modele(
        "Nettoyé", _pdf_standard(), "bulletin.pdf", correspondances=analyse.propositions
    )
    texte = pdf_export.texte_pdf(modele.contenu())
    assert "NOM PRENOMS" not in texte and "17010" not in texte and "Seventeen" not in texte
    assert "BULLETIN DE SOLDE / PAYSLIP" in texte and "Statut:" in texte
    balise = modele_bulletin_service.importer_modele("Balisé", _pdf_a_balises(), "b.pdf")
    assert "{{" not in pdf_export.texte_pdf(balise.contenu())


def test_correspondances_corrigees_par_l_administrateur():
    """L'administrateur peut retirer ou changer une proposition avant d'enregistrer."""
    analyse = modele_bulletin_service.analyser_modele(_pdf_standard(), "bulletin.pdf")
    correspondances = {cle: balise for cle, balise in analyse.propositions.items() if balise != "DETTE"}
    _contenu, zones = modele_bulletin_service.construire_modele(_pdf_standard(), analyse, correspondances)
    assert "DETTE" not in {b for z in zones for b in z.balises()}


def test_correspondances_sans_nom_ou_net_refusees():
    analyse = modele_bulletin_service.analyser_modele(_pdf_standard(), "bulletin.pdf")
    sans_net = {cle: balise for cle, balise in analyse.propositions.items() if balise != "NET_A_PAYER"}
    with pytest.raises(ModeleBulletinError, match="obligatoire"):
        modele_bulletin_service.importer_modele("X", _pdf_standard(), "b.pdf", correspondances=sans_net)


# ---------------------------------------------------------------------
# Refus de fichiers
# ---------------------------------------------------------------------

@pytest.mark.parametrize("nom, contenu, message", [
    ("bulletin.xlsx", b"PK\x03\x04...", "Format non accepté"),
    ("bulletin.pdf", b"PK\x03\x04 pas un pdf", "ne correspond pas"),
    ("bulletin.docx", b"%PDF-1.5 pas un docx", "ne correspond pas"),
    ("bulletin.pdf", b"", "vide"),
])
def test_fichiers_invalides_refuses(nom, contenu, message):
    with pytest.raises(ModeleBulletinError, match=message):
        modele_bulletin_service.analyser_modele(contenu, nom)


def test_fichier_trop_volumineux_refuse():
    with pytest.raises(ModeleBulletinError, match="10 Mo"):
        modele_bulletin_service.analyser_modele(b"%PDF" + b"0" * (10 * 1024 * 1024 + 1), "gros.pdf")


def test_pdf_sans_texte_refuse_avec_explication():
    document = pymupdf.open()
    document.new_page().draw_rect(pymupdf.Rect(10, 10, 100, 100), fill=(0, 0, 0))
    with pytest.raises(ModeleBulletinError, match="numérisé"):
        modele_bulletin_service.analyser_modele(document.tobytes(), "scan.pdf")


def test_balise_coupee_signalee():
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((50, 50), "{{NOM_COMPLET}} {{NET_A_", fontsize=11)
    page.insert_text((50, 80), "PAYER}}", fontsize=11)
    analyse = modele_bulletin_service.analyser_modele(document.tobytes(), "coupe.pdf")
    assert any("coupée" in erreur for erreur in analyse.erreurs)


# ---------------------------------------------------------------------
# Activation, suppression, repli
# ---------------------------------------------------------------------

def test_modele_actif_par_defaut_word_standard():
    assert modele_bulletin_service.obtenir_modele_actif().cle == modele_bulletin_service.CLE_STANDARD_WORD
    cles = [m.cle for m in modele_bulletin_service.lister_modeles()]
    assert cles[:2] == [modele_bulletin_service.CLE_STANDARD_WORD, modele_bulletin_service.CLE_STANDARD_PDF]


def test_suppression_protegee_puis_effective():
    modele = modele_bulletin_service.importer_modele("A", _docx_a_balises(), "a.docx", activer=True)
    with pytest.raises(ModeleBulletinError, match="actif"):
        modele_bulletin_service.supprimer_modele(modele.cle)
    with pytest.raises(ModeleBulletinError, match="standard"):
        modele_bulletin_service.supprimer_modele(modele_bulletin_service.CLE_STANDARD_PDF)
    modele_bulletin_service.activer_modele(modele_bulletin_service.CLE_STANDARD_WORD)
    modele_bulletin_service.supprimer_modele(modele.cle)
    assert not modele.chemin.exists()
    assert modele.cle not in [m.cle for m in modele_bulletin_service.lister_modeles()]


def test_repli_sur_le_modele_standard_si_fichier_disparu():
    modele = modele_bulletin_service.importer_modele("A", _pdf_a_balises(), "a.pdf", activer=True)
    modele.chemin.unlink()
    assert modele_bulletin_service.obtenir_modele_actif().cle == modele_bulletin_service.CLE_STANDARD_WORD
    e, p = _periode_validee_avec_enseignant()
    assert bulletin_service.generer_bulletin_enseignant(p.id, e.id).chemin.suffix == ".docx"


def test_aucune_alerte_quand_le_modele_choisi_est_utilisable():
    assert modele_bulletin_service.probleme_modele_actif() is None
    modele_bulletin_service.importer_modele("A", _docx_a_balises(), "a.docx", activer=True)
    assert modele_bulletin_service.probleme_modele_actif() is None


def test_alerte_a_l_ecran_si_le_fichier_du_modele_a_disparu():
    modele = modele_bulletin_service.importer_modele("Bulletin 2026", _pdf_a_balises(), "a.pdf", activer=True)
    modele.chemin.unlink()
    alerte = modele_bulletin_service.probleme_modele_actif()
    assert "« Bulletin 2026 »" in alerte and "introuvable" in alerte
    assert "modèle Word standard" in alerte and "Administration › Modèles de bulletin" in alerte


def test_alerte_a_l_ecran_si_le_modele_choisi_n_existe_plus():
    from database.repositories import parametres_paie_repository

    parametres_paie_repository.ecrire(modele_bulletin_service.CLE_MODELE_ACTIF, "importe_999")
    assert "n'existe plus" in modele_bulletin_service.probleme_modele_actif()
    assert modele_bulletin_service.obtenir_modele_actif().cle == modele_bulletin_service.CLE_STANDARD_WORD


def test_les_pages_qui_produisent_des_bulletins_affichent_l_alerte():
    for page in ("6_Bulletins_Paie.py", "17_Automatisation.py", "10_Administration.py"):
        assert "probleme_modele_actif()" in (RACINE / "ui_pages" / page).read_text(encoding="utf-8"), page


def test_regeneration_refusee_pour_periode_cloturee_quel_que_soit_le_format():
    modele_bulletin_service.activer_modele(modele_bulletin_service.CLE_STANDARD_PDF)
    e, p = _periode_validee_avec_enseignant()
    bulletin_service.generer_bulletin_enseignant(p.id, e.id)
    periode_service.cloturer_periode(p.id)
    modele_bulletin_service.activer_modele(modele_bulletin_service.CLE_STANDARD_WORD)
    with pytest.raises(bulletin_service.BulletinServiceError, match="clôturée"):
        bulletin_service.generer_bulletin_enseignant(p.id, e.id)


def test_import_et_activation_journalises(db_path):
    modele_bulletin_service.importer_modele("Journalisé", _docx_a_balises(), "j.docx", utilisateur="admin", activer=True)
    conn = sqlite3.connect(db_path)
    types = {l[0] for l in conn.execute("SELECT type_action FROM audit_log WHERE entite = 'modele_bulletin'")}
    conn.close()
    assert {TypeActionAudit.MODELE_BULLETIN_IMPORTE.value, TypeActionAudit.MODELE_BULLETIN_ACTIVE.value} <= types


def test_aucune_trace_si_import_echoue(tmp_path):
    with pytest.raises(ModeleBulletinError):
        modele_bulletin_service.importer_modele("X", _docx_a_balises(("PERIODE",)), "m.docx")
    dossier = tmp_path / "modeles"
    assert not dossier.exists() or not list(dossier.iterdir())
    assert len(modele_bulletin_service.lister_modeles()) == 2


# ---------------------------------------------------------------------
# Réservé à l'administrateur (côté service) — et conservé lors d'une réinitialisation
# ---------------------------------------------------------------------

@pytest.fixture
def comptes(db_path):
    return {
        role: utilisateur_service.creer_utilisateur("Nom", "Prenom", role.value, MDP, role, db_path=db_path)
        for role in RoleUtilisateur
    }


@pytest.mark.parametrize("role", [RoleUtilisateur.GESTIONNAIRE_PAIE, RoleUtilisateur.CONSULTATION])
def test_non_admin_refuse_meme_en_appel_direct(comptes, role, db_path):
    acteur = comptes[role].id
    with pytest.raises(AutorisationRefuseeError):
        administration_service.analyser_modele_bulletin(acteur, _docx_a_balises(), "m.docx", db_path=db_path)
    with pytest.raises(AutorisationRefuseeError):
        administration_service.importer_modele_bulletin(acteur, "X", _docx_a_balises(), "m.docx", db_path=db_path)
    with pytest.raises(AutorisationRefuseeError):
        administration_service.activer_modele_bulletin(acteur, modele_bulletin_service.CLE_STANDARD_PDF, db_path=db_path)
    with pytest.raises(AutorisationRefuseeError):
        administration_service.supprimer_modele_bulletin(acteur, "importe_1", db_path=db_path)
    assert modele_bulletin_service.cle_modele_actif(db_path=db_path) == modele_bulletin_service.CLE_STANDARD_WORD


def test_admin_autorise(comptes, db_path):
    admin = comptes[RoleUtilisateur.ADMIN].id
    modele = administration_service.importer_modele_bulletin(
        admin, "Admin", _docx_a_balises(), "m.docx", activer=True, db_path=db_path
    )
    assert modele_bulletin_service.cle_modele_actif(db_path=db_path) == modele.cle


FONCTIONS_SENSIBLES = {"importer_modele", "activer_modele", "supprimer_modele", "analyser_modele",
                       "definir_taux_taxe_defaut", "definir_taux_taxe_periode"}


@pytest.mark.parametrize("page", sorted((RACINE / "ui_pages").glob("*.py")), ids=lambda p: p.name)
def test_pages_passent_par_la_facade_administration(page):
    for noeud in ast.walk(ast.parse(page.read_text(encoding="utf-8"))):
        if isinstance(noeud, ast.Call):
            fonction = noeud.func
            nom = fonction.attr if isinstance(fonction, ast.Attribute) else getattr(fonction, "id", None)
            module = fonction.value.id if isinstance(fonction, ast.Attribute) and isinstance(fonction.value, ast.Name) else None
            if module != "administration_service":
                assert nom not in FONCTIONS_SENSIBLES, f"{page.name} appelle {module}.{nom} directement"


def test_reinitialisation_conserve_modeles_et_taux(comptes, db_path, tmp_path):
    from services import parametres_paie_service, reinitialisation_service

    admin = comptes[RoleUtilisateur.ADMIN]
    modele = modele_bulletin_service.importer_modele("Conservé", _pdf_a_balises(), "c.pdf", activer=True, db_path=db_path)
    parametres_paie_service.definir_taux_taxe_defaut("5.5", db_path=db_path)
    _periode_validee_avec_enseignant()
    rapport = reinitialisation_service.reinitialiser_donnees_metier(
        admin.id, "RÉINITIALISER", confirmation=True, db_path=db_path,
        dossier_exports=tmp_path / "exports", backup_dir=tmp_path / "sauvegardes",
    )
    assert rapport.reussie, rapport.message
    assert modele.chemin.exists()
    assert modele_bulletin_service.cle_modele_actif(db_path=db_path) == modele.cle
    assert str(parametres_paie_service.obtenir_taux_taxe_defaut(db_path=db_path)) == "0.055"
    assert periode_service.lister_periodes(db_path=db_path) == []


def test_catalogue_des_balises_coherent():
    noms = [b.nom for b in BALISES]
    assert len(noms) == len(set(noms))
    assert set(valeurs_exemple()) >= set(noms)
