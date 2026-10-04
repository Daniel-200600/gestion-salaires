"""
Tests de exports/reporting_export.py (module 13).
"""

import pytest
from openpyxl import load_workbook

import database.connection as database_connection
from services import (
    donnees_paie_service,
    enseignant_service,
    historique_paie_service,
    periode_service,
    reporting_paie_service,
)
from services.donnees_paie_service import DonneesPaieEnseignant

# Cas de référence officiel : vacataire, taxe de 5 %.
pytestmark = pytest.mark.taxe_historique


@pytest.fixture(autouse=True)
def _rediriger_connexion_par_defaut(db_path, monkeypatch, tmp_path):
    monkeypatch.setattr(database_connection, "DB_PATH", db_path)
    from exports import excel_export, reporting_export
    monkeypatch.setattr(excel_export, "EXPORT_DIR", tmp_path / "exports")
    monkeypatch.setattr(reporting_export, "EXPORT_DIR", tmp_path / "exports")


def _creer_enseignant(**overrides):
    donnees = {"nom": "Kamgang", "prenom": "Jean Paul", "sexe": "M", "statut": "V", "taux_horaire": 2000}
    donnees.update(overrides)
    return enseignant_service.creer_enseignant(**donnees)


def _creer_periode_ouverte(mois=1, annee=2051):
    p = periode_service.creer_periode(mois=mois, annee=annee)
    return periode_service.ouvrir_periode(p.id)


def _saisir_donnees(periode_id, enseignant_id, **kwargs):
    donnees_paie_service.enregistrer_donnees_paie_groupe(
        periode_id, [DonneesPaieEnseignant(enseignant_id=enseignant_id, **kwargs)]
    )


def _rapport_cas_reference():
    e = _creer_enseignant(taux_horaire=2000)
    p = _creer_periode_ouverte()
    _saisir_donnees(
        p.id, e.id, heures_par_semaine={1: 20, 2: 20, 3: 20, 4: 20, 5: 20},
        prime_ap_pp=20000, surveillance_secretariat=10000, indemnite_suggestion_admin=5000,
        retenue_amicale=5000, dette=10000,
    )
    return reporting_paie_service.construire_etat_paie_complet(p.id, etablissement="Ecole Test", utilisateur="admin")


# ---------------------------------------------------------------------
# Export complet
# ---------------------------------------------------------------------

def test_export_cree_fichier():
    from exports.reporting_export import generer_fichier_etat_paie
    rapport = _rapport_cas_reference()
    chemin = generer_fichier_etat_paie(rapport)
    assert chemin.exists()
    assert chemin.suffix == ".xlsx"


def test_export_nom_fichier():
    from exports.reporting_export import generer_fichier_etat_paie
    rapport = _rapport_cas_reference()
    chemin = generer_fichier_etat_paie(rapport)
    assert chemin.name.startswith("Etat_Paie_")


def test_export_feuilles_attendues_sans_comparaison():
    from exports.reporting_export import generer_fichier_etat_paie
    rapport = _rapport_cas_reference()
    chemin = generer_fichier_etat_paie(rapport)
    wb = load_workbook(chemin)
    assert wb.sheetnames == ["Synthese", "Detail_Enseignants", "Par_Statut", "Par_Sexe", "Retenues", "Rapprochement"]


def test_export_feuille_comparaison_presente_si_fournie():
    from exports.reporting_export import generer_fichier_etat_paie

    e = _creer_enseignant()
    p1 = _creer_periode_ouverte(mois=2)
    _saisir_donnees(p1.id, e.id, heures_par_semaine={1: 10})
    rapport = _rapport_cas_reference()

    comparaison = historique_paie_service.comparer_periodes(p1.id, rapport.periode.id)
    chemin = generer_fichier_etat_paie(rapport, comparaison=comparaison)
    wb = load_workbook(chemin)
    assert "Comparaison" in wb.sheetnames


def test_export_ne_ecrase_jamais_un_fichier_existant():
    from exports.reporting_export import generer_fichier_etat_paie
    rapport = _rapport_cas_reference()
    chemin1 = generer_fichier_etat_paie(rapport)
    chemin2 = generer_fichier_etat_paie(rapport)
    assert chemin1 != chemin2
    assert chemin1.exists() and chemin2.exists()


# ---------------------------------------------------------------------
# Contenu / totaux
# ---------------------------------------------------------------------

def test_export_synthese_contient_le_net_officiel():
    from exports.reporting_export import generer_fichier_etat_paie
    rapport = _rapport_cas_reference()
    chemin = generer_fichier_etat_paie(rapport)
    wb = load_workbook(chemin)
    valeurs = [c.value for c in wb["Synthese"]["B"] if c.value is not None]
    assert 208250 in valeurs


def test_export_detail_enseignants_contient_toutes_les_colonnes():
    from exports.reporting_export import generer_fichier_etat_paie
    rapport = _rapport_cas_reference()
    chemin = generer_fichier_etat_paie(rapport)
    wb = load_workbook(chemin)
    ws = wb["Detail_Enseignants"]
    entetes = [c.value for c in ws[3]]  # ligne d'en-têtes (après titre + ligne vide)
    for colonne_attendue in ["Nom", "Prénom", "Sexe", "Statut", "Net à payer"]:
        assert colonne_attendue in entetes


def test_export_rapprochement_toutes_lignes_ok():
    from exports.reporting_export import generer_fichier_etat_paie
    rapport = _rapport_cas_reference()
    chemin = generer_fichier_etat_paie(rapport)
    wb = load_workbook(chemin)
    ws = wb["Rapprochement"]
    statuts = [row[4].value for row in ws.iter_rows(min_row=4) if row[4].value]
    assert len(statuts) > 0
    assert all(s == "OK" for s in statuts)


def test_export_retenues_totaux_presents():
    from exports.reporting_export import generer_fichier_etat_paie
    rapport = _rapport_cas_reference()
    chemin = generer_fichier_etat_paie(rapport)
    wb = load_workbook(chemin)
    ws = wb["Retenues"]
    valeurs_col_a = [c.value for c in ws["A"] if c.value == "TOTAL"]
    assert len(valeurs_col_a) == 1
