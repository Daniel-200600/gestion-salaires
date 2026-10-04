"""
Tests de exports/statistiques_export.py et des permissions du module 17.
"""

import pytest
from openpyxl import load_workbook

import database.connection as database_connection
from models.enums import RoleUtilisateur
from services import controle_paie_service, donnees_paie_service, enseignant_service, periode_service, permission_service, reporting_paie_service, statistiques_service
from services.comptabilite_service import preparer_etat_comptable
from services.donnees_paie_service import DonneesPaieEnseignant

# Cas de référence officiel : vacataire, taxe de 5 %.
pytestmark = pytest.mark.taxe_historique


@pytest.fixture(autouse=True)
def _rediriger_connexion_par_defaut(db_path, monkeypatch, tmp_path):
    monkeypatch.setattr(database_connection, "DB_PATH", db_path)
    from exports import excel_export, statistiques_export
    monkeypatch.setattr(excel_export, "EXPORT_DIR", tmp_path / "exports")
    monkeypatch.setattr(statistiques_export, "EXPORT_DIR", tmp_path / "exports")


def _creer_enseignant(**overrides):
    donnees = {"nom": "Kamgang", "prenom": "Jean Paul", "sexe": "M", "statut": "V", "taux_horaire": 2000}
    donnees.update(overrides)
    return enseignant_service.creer_enseignant(**donnees)


def _rapport_cas_reference():
    e = _creer_enseignant()
    p = periode_service.creer_periode(mois=1, annee=2098)
    controle_paie_service.ouvrir_periode_avec_audit(p.id)
    donnees_paie_service.enregistrer_donnees_paie_groupe(p.id, [
        DonneesPaieEnseignant(enseignant_id=e.id, heures_par_semaine={1: 20, 2: 20, 3: 20, 4: 20, 5: 20},
            prime_ap_pp=20000, surveillance_secretariat=10000, indemnite_suggestion_admin=5000,
            retenue_amicale=5000, dette=10000),
    ])
    etat = preparer_etat_comptable(p.id)
    stats = statistiques_service.statistiques_generales(p.id)
    par_statut = reporting_paie_service.synthese_detaillee_par_statut(etat.resultats)
    par_sexe = reporting_paie_service.synthese_detaillee_par_sexe(etat.resultats)
    composantes = statistiques_service.analyser_composantes(etat.resultats)
    return stats, etat.resultats, par_statut, par_sexe, composantes


# ---------------------------------------------------------------------
# Export — présence des feuilles et cohérence
# ---------------------------------------------------------------------

def test_export_cree_fichier():
    from exports.statistiques_export import generer_fichier_statistiques
    stats, resultats, par_statut, par_sexe, composantes = _rapport_cas_reference()
    chemin = generer_fichier_statistiques(stats, resultats, par_statut, par_sexe, composantes)
    assert chemin.exists()
    assert chemin.name.startswith("Statistiques_")


def test_export_feuilles_attendues():
    from exports.statistiques_export import generer_fichier_statistiques
    stats, resultats, par_statut, par_sexe, composantes = _rapport_cas_reference()
    chemin = generer_fichier_statistiques(stats, resultats, par_statut, par_sexe, composantes)
    wb = load_workbook(chemin)
    for feuille_attendue in ["Synthese", "Enseignants", "Par_Statut", "Par_Sexe", "Composantes"]:
        assert feuille_attendue in wb.sheetnames


def test_export_cas_reference_coherent():
    from exports.statistiques_export import generer_fichier_statistiques
    stats, resultats, par_statut, par_sexe, composantes = _rapport_cas_reference()
    chemin = generer_fichier_statistiques(stats, resultats, par_statut, par_sexe, composantes)
    wb = load_workbook(chemin)
    valeurs = [c.value for c in wb["Synthese"]["B"] if c.value is not None]
    assert 208250 in valeurs


def test_export_avec_outliers_et_periodes():
    from exports.statistiques_export import generer_fichier_statistiques
    stats, resultats, par_statut, par_sexe, composantes = _rapport_cas_reference()
    outliers = statistiques_service.detecter_valeurs_atypiques(resultats, "taux_horaire")
    points = statistiques_service.analyser_periodes([stats.periode.id])
    chemin = generer_fichier_statistiques(stats, resultats, par_statut, par_sexe, composantes, outliers=outliers, points_periodes=points)
    wb = load_workbook(chemin)
    assert "Outliers" in wb.sheetnames
    assert "Par_Periode" in wb.sheetnames


def test_export_ne_ecrase_jamais_un_fichier_existant():
    from exports.statistiques_export import generer_fichier_statistiques
    stats, resultats, par_statut, par_sexe, composantes = _rapport_cas_reference()
    chemin1 = generer_fichier_statistiques(stats, resultats, par_statut, par_sexe, composantes)
    chemin2 = generer_fichier_statistiques(stats, resultats, par_statut, par_sexe, composantes)
    assert chemin1 != chemin2
    assert chemin1.exists() and chemin2.exists()


# ---------------------------------------------------------------------
# Permissions (section 22)
# ---------------------------------------------------------------------

def test_permission_statistiques_consulter_accessible_aux_trois_roles():
    assert permission_service.a_permission(RoleUtilisateur.ADMIN, permission_service.STATISTIQUES_CONSULTER)
    assert permission_service.a_permission(RoleUtilisateur.GESTIONNAIRE_PAIE, permission_service.STATISTIQUES_CONSULTER)
    assert permission_service.a_permission(RoleUtilisateur.CONSULTATION, permission_service.STATISTIQUES_CONSULTER)


def test_permission_statistiques_exporter_refusee_a_consultation():
    assert permission_service.a_permission(RoleUtilisateur.ADMIN, permission_service.STATISTIQUES_EXPORTER)
    assert permission_service.a_permission(RoleUtilisateur.GESTIONNAIRE_PAIE, permission_service.STATISTIQUES_EXPORTER)
    assert not permission_service.a_permission(RoleUtilisateur.CONSULTATION, permission_service.STATISTIQUES_EXPORTER)
