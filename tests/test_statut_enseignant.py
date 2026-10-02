"""
Changement de statut d'un enseignant (Vacataire <-> Permanent).

Le nouveau statut s'applique aux calculs et bulletins produits ensuite ;
un bulletin déjà émis pour une période validée ou clôturée garde son
statut ; chaque changement est journalisé ; la page Enseignants propose
une action dédiée.
"""

import sqlite3
from pathlib import Path

import pytest
from docx import Document

import database.connection as database_connection
from exports.word_export import texte_document
from models.enums import StatutEnseignant, TypeActionAudit
from services import bulletin_service, donnees_paie_service, enseignant_service, paie_service, periode_service
from services.donnees_paie_service import DonneesPaieEnseignant
from services.enseignant_service import EnseignantValidationError

RACINE = Path(__file__).resolve().parent.parent


@pytest.fixture(autouse=True)
def _base(db_path, monkeypatch, tmp_path):
    monkeypatch.setattr(database_connection, "DB_PATH", db_path)
    monkeypatch.setattr(bulletin_service, "EXPORT_DIR_BULLETINS", tmp_path / "bulletins")


def _enseignant(statut="V"):
    return enseignant_service.creer_enseignant(nom="Fotso", prenom="Paul", sexe="M", statut=statut, taux_horaire=1500)


def _periode_avec_heures(enseignant_id, mois):
    p = periode_service.ouvrir_periode(periode_service.creer_periode(mois=mois, annee=2026).id)
    donnees_paie_service.enregistrer_donnees_paie_groupe(
        p.id, [DonneesPaieEnseignant(enseignant_id=enseignant_id, heures_par_semaine={1: 10})]
    )
    return p


def _audits(db_path):
    conn = sqlite3.connect(db_path)
    lignes = conn.execute(
        "SELECT entite_id, utilisateur, details FROM audit_log WHERE type_action = ?",
        (TypeActionAudit.STATUT_ENSEIGNANT_MODIFIE.value,),
    ).fetchall()
    conn.close()
    return lignes


def test_changer_statut_vacataire_vers_permanent_et_retour(db_path):
    e = _enseignant("V")
    e = enseignant_service.changer_statut_enseignant(e.id, "P", utilisateur="gestion")
    assert e.statut == StatutEnseignant.PERMANENT
    e = enseignant_service.changer_statut_enseignant(e.id, StatutEnseignant.VACATAIRE, utilisateur="gestion")
    assert e.statut == StatutEnseignant.VACATAIRE
    assert _audits(db_path) == [
        (e.id, "gestion", "Fotso Paul : Vacataire -> Permanent"),
        (e.id, "gestion", "Fotso Paul : Permanent -> Vacataire"),
    ]


def test_meme_statut_refuse():
    e = _enseignant("P")
    with pytest.raises(EnseignantValidationError, match="déjà"):
        enseignant_service.changer_statut_enseignant(e.id, "P")


def test_statut_invalide_ou_enseignant_inexistant_refuse():
    e = _enseignant("V")
    with pytest.raises(EnseignantValidationError):
        enseignant_service.changer_statut_enseignant(e.id, "X")
    with pytest.raises(EnseignantValidationError):
        enseignant_service.changer_statut_enseignant(9999, "P")


def test_changement_via_formulaire_de_modification_journalise(db_path):
    e = _enseignant("V")
    enseignant_service.modifier_enseignant(e.id, "Fotso", "Paul", "M", "P", 1500, utilisateur="admin")
    assert len(_audits(db_path)) == 1
    enseignant_service.modifier_enseignant(e.id, "Fotso", "Paul", "M", "P", 1600, utilisateur="admin")
    assert len(_audits(db_path)) == 1  # pas de changement de statut : pas de nouvelle entrée


def test_nouveau_statut_applique_aux_periodes_en_cours():
    e = _enseignant("V")
    p = _periode_avec_heures(e.id, mois=3)
    enseignant_service.changer_statut_enseignant(e.id, "P")
    assert paie_service.calculer_paie_enseignant(p.id, e.id).statut == StatutEnseignant.PERMANENT
    p = periode_service.valider_periode(p.id)
    texte = texte_document(Document(bulletin_service.generer_bulletin_enseignant(p.id, e.id).chemin))
    assert "Statut: P" in texte


def test_bulletin_deja_emis_conserve_son_statut():
    e = _enseignant("V")
    p = periode_service.valider_periode(_periode_avec_heures(e.id, mois=4).id)
    premier = bulletin_service.generer_bulletin_enseignant(p.id, e.id)  # instantané : statut V
    enseignant_service.changer_statut_enseignant(e.id, "P")
    second = bulletin_service.generer_bulletin_enseignant(p.id, e.id)   # période validée : régénération
    assert "Statut: V" in texte_document(Document(second.chemin))
    assert second.resultat.statut == StatutEnseignant.VACATAIRE
    assert premier.chemin != second.chemin

    # Une nouvelle période utilise le nouveau statut.
    p2 = periode_service.valider_periode(_periode_avec_heures(e.id, mois=5).id)
    assert "Statut: P" in texte_document(Document(bulletin_service.generer_bulletin_enseignant(p2.id, e.id).chemin))


def test_page_enseignants_propose_l_action_dediee():
    source = (RACINE / "ui_pages" / "1_Enseignants.py").read_text(encoding="utf-8")
    assert '"Changer le statut"' in source
    assert "changer_statut_enseignant(" in source
    assert "ENSEIGNANT_MODIFIER" in source
