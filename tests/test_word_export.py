"""
Tests du générateur Word (exports/word_export.py).

Vérifie le chargement du template, le remplacement des placeholders,
la détection de placeholders oubliés, le nommage de fichier, et
l'absence de toute formule de paie dans ce module.
"""

import inspect

import pytest
from docx import Document

from exports import word_export
from exports.word_export import (
    TEMPLATE_PATH,
    WordExportError,
    generer_document_bulletin,
    generer_nom_fichier_bulletin,
    sauvegarder_document,
    verifier_aucun_placeholder_restant,
)


VALEURS_COMPLETES = {
    "{{NOM}}": "KAMGANG",
    "{{PRENOM}}": "JEAN PAUL",
    "{{STATUT}}": "P",
    "{{PERIODE}}": "AOUT 2026",
    "{{TOTAL_HEURES}}": "100",
    "{{TAUX_HORAIRE}}": "2000",
    "{{GAIN_HEURES}}": "200000",
    "{{PRIME_AP_PP}}": "20000",
    "{{SURVEILLANCE_SECRETARIAT}}": "10000",
    "{{INDEMNITE_SUGGESTION_ADMIN}}": "5000",
    "{{TAXE_5}}": "11750",
    "{{RETENUE_AMICALE}}": "5000",
    "{{DETTE}}": "10000",
    "{{TOTAL_GAINS}}": "235000",
    "{{TOTAL_RETENUES}}": "26750",
    "{{NET_A_PERÇEVOIR}}": "208250",
    "{{NET_EN_LETTRES}}": "Two Hundred Eight Thousand Two Hundred Fifty",
    "{{DATE_GENERATION}}": "31/08/2026",
}


def _lire_texte_document(document) -> str:
    textes = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                textes.append(cell.text)
    return " | ".join(textes)


# ---------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------

def test_template_existe():
    assert TEMPLATE_PATH.exists(), "Le template doit être généré une fois via templates/build_template.py"


def test_template_absent_leve_erreur_claire(tmp_path):
    faux_chemin = tmp_path / "inexistant.docx"
    with pytest.raises(WordExportError, match="introuvable"):
        generer_document_bulletin(VALEURS_COMPLETES, template_path=faux_chemin)


# ---------------------------------------------------------------------
# Remplacement des placeholders
# ---------------------------------------------------------------------

def test_generation_document_remplace_tous_les_placeholders():
    document = generer_document_bulletin(VALEURS_COMPLETES)
    texte = _lire_texte_document(document)
    assert "{{" not in texte
    assert "}}" not in texte


def test_valeurs_correctement_injectees():
    document = generer_document_bulletin(VALEURS_COMPLETES)
    texte = _lire_texte_document(document)
    assert "KAMGANG" in texte
    assert "JEAN PAUL" in texte
    assert "208250" in texte
    assert "Two Hundred Eight Thousand Two Hundred Fifty" in texte


def test_verifier_aucun_placeholder_leve_si_oubli():
    valeurs_incompletes = dict(VALEURS_COMPLETES)
    del valeurs_incompletes["{{NET_A_PERÇEVOIR}}"]  # un placeholder volontairement non fourni

    document = Document(TEMPLATE_PATH)
    from exports.word_export import _remplacer_dans_document
    _remplacer_dans_document(document, valeurs_incompletes)

    with pytest.raises(WordExportError, match="[Pp]laceholder"):
        verifier_aucun_placeholder_restant(document)


# ---------------------------------------------------------------------
# Nom de fichier
# ---------------------------------------------------------------------

def test_nom_fichier_format_attendu():
    nom = generer_nom_fichier_bulletin("Kamgang", "Jean Paul", "Août 2026")
    assert nom == "Bulletin_KAMGANG_JEAN_PAUL_AOÛT_2026.docx"


def test_nom_fichier_caracteres_interdits_nettoyes():
    nom = generer_nom_fichier_bulletin('Test:<>', 'Nom"Bizarre/\\|?*', "Période")
    for caractere_interdit in '<>:"/\\|?*':
        assert caractere_interdit not in nom


def test_nom_fichier_ne_contient_jamais_generique():
    nom = generer_nom_fichier_bulletin("Ngono", "Marie", "Septembre 2026")
    assert nom != "bulletin.docx"
    assert "NGONO" in nom


# ---------------------------------------------------------------------
# Sauvegarde
# ---------------------------------------------------------------------

def test_sauvegarde_document(tmp_path):
    document = generer_document_bulletin(VALEURS_COMPLETES)
    chemin = tmp_path / "sortie" / "test.docx"
    resultat = sauvegarder_document(document, chemin)
    assert resultat == chemin
    assert chemin.exists()
    Document(chemin)  # relecture valide


# ---------------------------------------------------------------------
# Aucune formule de paie dans word_export.py (règle absolue du module 07)
# ---------------------------------------------------------------------

def test_aucune_formule_de_paie():
    source = inspect.getsource(word_export)
    assert "* 0.05" not in source
    assert "taux_horaire *" not in source
    assert "gain_heures =" not in source
    assert "net_a_percevoir =" not in source
